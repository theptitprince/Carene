# -*- coding: utf-8 -*-
"""Le « fichier type à remplir » : UN classeur Excel pour tout le navire.

Jusqu'ici, reconstituer un navire demandait de fabriquer soi-même un fichier
par table, avec des intitulés de colonnes à deviner. Ce module donne le chemin
court :

- `ecrire_classeur(chemin)` écrit un classeur VIDE — une feuille par table, la
  ligne d'intitulés canoniques déjà en place, et sur chaque intitulé un
  COMMENTAIRE de cellule qui dit l'unité, le repère, si c'est obligatoire, et
  la valeur du navire d'exemple. Ce sont exactement les explications de la
  fenêtre « Création du navire » (`ship_editor_content.format_rows`), amenées
  là où l'utilisateur travaille vraiment : dans son tableur ;
- `importer_classeur(draft, chemin)` relit toutes les feuilles d'un classeur
  rempli, en un seul geste, et rend un compte rendu feuille par feuille ;
- `ecrire_modeles_csv(dossier)` produit les mêmes gabarits en CSV, pour qui n'a
  pas Excel.

Rien du format du dossier navire n'est redit ici : les colonnes viennent des
`SPEC_*` de `navire_draft`, les explications de `ship_editor_content`. Ajouter
une colonne à une table la fait apparaître dans le classeur toute seule.
"""
from __future__ import annotations

import csv
import os
import re

from .navire_draft import (
    CHAMPS_DIMENSIONS,
    CHAMPS_IDENTIFICATION,
    CHAMPS_LEGE,
    CHAMPS_VENT,
    SPEC_CAPACITES,
    SPEC_CAS_REFERENCE,
    SPEC_COUPLES,
    SPEC_ENVAHISSEMENT,
    SPEC_HYDRO,
    SPEC_JAUGE,
    SPEC_PROFILS_VENT,
    normalize,
    table_de_feuille,
    to_float,
)

# Angles de gîte du gabarit : ceux du dossier de référence. L'utilisateur ajoute ou
# retire des colonnes selon sa propre table — l'import lit l'angle dans
# l'intitulé, il n'y a aucune liste figée côté lecture.
ANGLES_KN_TYPE = [5, 10, 15, 20, 25, 30, 40, 50, 60]

# Noms de feuilles, sans accent (un accent dans un nom de feuille se perd d'un
# tableur à l'autre). La reconnaissance à l'import passe par `normalize` :
# « Pantocarènes KN » et « pantocarenes_kn » désignent la même feuille.
F_LISEZMOI = "Lisez-moi"
F_IDENTIFICATION = "Identification"
F_LEGE = "Lege"
F_VENT = "Vent"
F_PROFILS_VENT = "Profils_vent"
F_COUPLES = "Couples"
F_HYDRO = "Hydrostatiques"
F_KN = "Pantocarenes_KN"
F_CAPACITES = "Capacites"
F_JAUGE_EXEMPLE = "Jauge_EXEMPLE"
F_ENVAHISSEMENT = "Envahissement"
F_CAS = "Cas_de_reference"

PREFIXE_JAUGE = "Jauge_"

# à quoi sert chaque feuille : repris tel quel dans « Lisez-moi » et dans le
# texte de l'étape « classeur » de la fenêtre
FEUILLES_DESCRIPTION = [
    (F_IDENTIFICATION, "Nom, immatriculation et dimensions principales. Une "
                       "ligne par renseignement : la valeur va dans la "
                       "colonne « valeur »."),
    (F_LEGE, "Le navire vide : masse et centre de gravité (essai de "
             "stabilité ou page « lightship » du dossier)."),
    (F_VENT, "Surface exposée au vent d'une condition de référence. "
             "Facultatif : sans elle, seul le critère météo n'est pas évalué."),
    (F_PROFILS_VENT, "Les surfaces au vent de CHAQUE configuration de voilure "
                     "(cargo, voilure complète, intermédiaire, réduite…), "
                     "tirant d'eau par tirant d'eau. Facultatif, et réservé "
                     "aux navires à voiles : c'est elle qui permet d'évaluer "
                     "le critère de vent dans la voilure réellement portée."),
    (F_COUPLES, "La table d'espacement des couples (n, x). Sert à caler les "
                "plans « sur le couple 35 » ; ne sert pas au calcul."),
    (F_HYDRO, "La table hydrostatique, TOUTES assiettes dans la même feuille : "
              "la colonne « Assiette_m » dit à laquelle appartient chaque "
              "ligne."),
    (F_KN, "Les pantocarènes : une ligne par (assiette, déplacement), une "
           "colonne par angle de gîte."),
    (F_CAPACITES, "La liste des capacités liquides : contenu, densité, volume, "
                  "FSM maximal."),
    (F_JAUGE_EXEMPLE, "Modèle de table de jaugeage. À DUPLIQUER : une feuille "
                      "« Jauge_<NOM DE LA CAPACITÉ> » par capacité (le nom "
                      "doit être celui de la feuille « Capacites » ; un "
                      "souligné y remplace un espace ou une barre oblique)."),
    (F_ENVAHISSEMENT, "Les ouvertures non étanches et leurs coordonnées."),
    (F_CAS, "Des cas de chargement du dossier approuvé avec leurs résultats "
            "imprimés : le logiciel les recalcule et compare. C'est le "
            "contrôle de tout le reste."),
]


# ------------------------------------------------------------- légendes
def _apparier(spec, format_rows):
    """{nom canonique: (unité, requis, exemple)} depuis la table « Format
    attendu » de la fenêtre.

    Les intitulés de cette table sont humains (« LCB — centre de carène
    long. ») ; les colonnes du classeur sont canoniques (`LCB_m`). On les
    rapproche avec les mêmes synonymes qu'à l'import, en préférant la
    correspondance exacte, puis le préfixe, puis l'alias le plus long : sans ce
    départage, « MCT — moment unitaire d'assiette » allait à `Assiette_m`.
    """
    scores = []
    for col in spec:
        for i, (nom, unite, requis, exemple) in enumerate(format_rows):
            hn = normalize(nom)
            meilleur = 0
            for a in col.aliases:
                if hn == a:
                    meilleur = max(meilleur, 1000)
                elif len(a) >= 3 and hn.startswith(a):
                    meilleur = max(meilleur, 100 + len(a))
                elif len(a) >= 3 and a in hn:
                    meilleur = max(meilleur, len(a))
            if meilleur:
                scores.append((meilleur, col.name, i))
    scores.sort(key=lambda t: -t[0])
    out, pris_col, pris_row = {}, set(), set()
    for _s, nom, i in scores:
        if nom in pris_col or i in pris_row:
            continue
        pris_col.add(nom)
        pris_row.add(i)
        out[nom] = (format_rows[i][1], format_rows[i][2], format_rows[i][3])
    return out


# Commentaires écrits à la main là où la table « Format attendu » de la
# fenêtre ne dit pas tout : ce qui est propre au CLASSEUR (une seule feuille
# pour toutes les assiettes) n'a pas d'équivalent dans une fenêtre où chaque
# assiette s'importait dans un fichier séparé.
COMMENTAIRES_PROPRES = {
    (F_HYDRO, "Assiette_m"):
        "Unité / repère : m, positive sur l'arrière.\n"
        "Renseignement : obligatoire.\n"
        "UNE SEULE feuille porte toutes les assiettes du dossier : c'est "
        "cette colonne qui dit à laquelle appartient la ligne. Recopiez une "
        "table complète par assiette, les unes sous les autres.\n"
        "Exemple : -0.5, 0, 0.5, 1, 1.5 (cinq tables de 25 tirants "
        "d'eau)",
    (F_CAPACITES, "Nom"):
        "Unité / repère : texte — le code du dossier.\n"
        "Renseignement : obligatoire.\n"
        "C'est ce nom qui relie la capacité à sa feuille de jaugeage "
        "« Jauge_<NOM> » : les deux doivent concorder, au souligné près.\n"
        "Exemple : WB 3B",
}


def _commentaire_colonne(col, legende):
    """Le texte du commentaire posé sur un intitulé de colonne."""
    if legende:
        unite, requis, exemple = legende
        lignes = [f"Unité / repère : {unite}", f"Renseignement : {requis}"]
        if exemple and exemple != "—":
            lignes.append(f"Exemple : {exemple}")
        if col.required and "oui" not in str(requis).lower():
            # la fenêtre dit « recommandé » là où l'import refuse la table
            # sans : c'est l'import qui fait foi, et le dire évite un fichier
            # rejeté sans que l'utilisateur comprenne pourquoi
            lignes.append("Obligatoire à l'import : sans cette colonne, la "
                          "table est refusée.")
    else:
        lignes = ["Unité / repère : voir « Conventions » du Lisez-moi.",
                  "Renseignement : "
                  + ("obligatoire" if col.required else "facultatif")]
    vus, acceptes = {normalize(col.name)}, [col.name]
    for a in col.aliases[1:]:
        if len(a) >= 3 and a not in vus:
            vus.add(a)
            acceptes.append(a)
    lignes.append("Intitulés acceptés : " + ", ".join(acceptes[:7]))
    return "\n".join(lignes)


def _legendes_champs(champs, format_rows):
    """{clé de champ: (unité, requis, exemple)} pour une feuille clé/valeur."""
    par_libelle = {normalize(nom): (unite, requis, ex)
                   for nom, unite, requis, ex in format_rows}
    out = {}
    for cle, libelle, unite, requis in champs:
        trouve = None
        ln = normalize(libelle)
        for k, v in par_libelle.items():
            if k == ln or (len(k) >= 4 and (ln.startswith(k) or k.startswith(ln))):
                trouve = v
                break
        out[cle] = trouve or (unite, "oui" if requis else "facultatif", "")
    return out


def _etapes():
    """Le contenu rédactionnel de la fenêtre, importé tardivement : le module
    de contenu n'a rien à voir avec le moteur, et un import en tête ferait
    dépendre `carene.core` de l'interface."""
    from ..ship_editor_content import CONVENTIONS, STEPS
    return {s["key"]: s for s in STEPS}, CONVENTIONS


# ------------------------------------------------------------- écriture
def _entetes_kn(angles):
    return ["Assiette_m", "Deplacement_t"] + [f"KN_{float(a):g}" for a in angles]


# Colonnes qui restent du TEXTE même quand elles ressemblent à un nombre : le
# cas « 01 » du dossier de référence deviendrait 1 en passant par un nombre, et le
# rejeu ne retrouverait plus son cas. Ce sont les mêmes colonnes que celles
# écartées de la conversion à l'import (`numeric_except`).
COLONNES_TEXTE = {"Nom", "Type", "Groupe", "Repere", "Fonction", "Cas",
                  "Titre", "Profil", "Source"}


def _valeur_cellule(v, texte=False):
    """Ce qu'on écrit vraiment dans une cellule : un nombre reste un nombre
    (sinon le tableur affiche des apostrophes et les tris se font en texte)."""
    if v is None:
        return ""
    if texte:
        return "" if v == "" else str(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "":
        return ""
    f = to_float(s)
    return f if f is not None and not re.search(r"[A-Za-z]", s) else s


def ecrire_classeur(chemin_xlsx, draft=None):
    """Écrit le classeur type. `draft` non nul : les feuilles sont remplies avec
    ses données (c'est ainsi qu'est produit le classeur d'exemple).

    Retourne la liste des noms de feuilles écrites.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as e:  # pragma: no cover
        raise RuntimeError(
            "L'écriture du classeur demande le module openpyxl "
            "(pip install openpyxl). Utilisez « Créer les gabarits CSV… » à "
            "la place.") from e

    etapes, conventions = _etapes()
    gras = Font(bold=True)

    wb = Workbook()
    wb.remove(wb.active)

    def feuille(nom):
        # 31 caractères : la limite d'Excel. Deux noms de capacité qui se
        # ressemblent sur 31 caractères donneraient la même feuille — on
        # numérote plutôt que d'en perdre une en silence.
        titre = nom[:31]
        n = 2
        while titre in wb.sheetnames:
            suffixe = f"~{n}"
            titre = nom[:31 - len(suffixe)] + suffixe
            n += 1
        return wb.create_sheet(titre)

    def entetes(ws, noms, commentaires=None, largeurs=None):
        for i, nom in enumerate(noms, start=1):
            c = ws.cell(row=1, column=i, value=nom)
            c.font = gras
            c.alignment = Alignment(horizontal="left")
            txt = (commentaires or {}).get(nom)
            if txt:
                # 320 × 180 px : de quoi lire quatre lignes sans redimensionner
                com = Comment(txt, "Carène")
                com.width, com.height = 320, 180
                c.comment = com
            largeur = (largeurs or {}).get(nom, max(12, len(nom) + 3))
            ws.column_dimensions[get_column_letter(i)].width = largeur
        # la ligne d'intitulés reste visible quand on descend dans la table
        ws.freeze_panes = "A2"

    def table(nom_feuille, spec, cle_etape, lignes=None, noms=None,
              cle_format="format_rows"):
        ws = feuille(nom_feuille)
        etape = etapes.get(cle_etape, {})
        legendes = _apparier(spec, etape.get(cle_format) or [])
        noms = noms or [c.name for c in spec]
        par_nom = {c.name: c for c in spec}
        commentaires = {}
        for n in noms:
            col = par_nom.get(n)
            if col is not None:
                commentaires[n] = COMMENTAIRES_PROPRES.get(
                    (nom_feuille, n)) or _commentaire_colonne(col,
                                                              legendes.get(n))
        entetes(ws, noms, commentaires)
        for r, ligne in enumerate(lignes or [], start=2):
            for i, n in enumerate(noms, start=1):
                ws.cell(row=r, column=i, value=_valeur_cellule(
                    ligne.get(n), texte=n in COLONNES_TEXTE))
        return ws

    def cle_valeur(nom_feuille, champs, valeurs, cle_etape):
        ws = feuille(nom_feuille)
        etape = etapes.get(cle_etape, {})
        legendes = _legendes_champs(champs, etape.get("format_rows") or [])
        entetes(ws, ["champ", "valeur", "libelle"], {
            "champ": "Le nom technique du renseignement : ne le modifiez pas, "
                     "c'est lui que le logiciel reconnaît.\nSurvolez chaque "
                     "ligne pour l'unité et un exemple.",
            "valeur": "C'est ici que vous écrivez. Laissez vide ce que votre "
                      "dossier ne donne pas.",
            "libelle": "Le même renseignement en clair. Colonne d'aide : elle "
                       "n'est pas relue.",
        }, largeurs={"champ": 32, "valeur": 18, "libelle": 46})
        for r, (cle, libelle, unite, requis) in enumerate(champs, start=2):
            c = ws.cell(row=r, column=1, value=cle)
            unite_l, requis_l, exemple = legendes.get(
                cle, (unite, "oui" if requis else "facultatif", ""))
            txt = [f"{libelle}", f"Unité / repère : {unite_l}",
                   f"Renseignement : {requis_l}"]
            if exemple and exemple != "—":
                txt.append(f"Exemple : {exemple}")
            com = Comment("\n".join(txt), "Carène")
            com.width, com.height = 320, 150
            c.comment = com
            # un champ textuel reste du texte : un numéro OMI à zéro initial
            # perdrait son zéro en passant par un nombre
            ws.cell(row=r, column=2,
                    value=_valeur_cellule((valeurs or {}).get(cle),
                                          texte=(unite == "texte")))
            ws.cell(row=r, column=3, value=libelle + (" *" if requis else ""))
        return ws

    # ------------------------------------------------------- Lisez-moi
    ws = feuille(F_LISEZMOI)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    lignes_lm = [("Classeur de reprise d'un dossier de stabilité", ""), ("", "")]
    lignes_lm.append(("À quoi sert ce classeur",
                      "Recopier, une fois pour toutes, les tables de votre "
                      "dossier de stabilité approuvé. Une feuille par table. "
                      "Remplissez ce que vous avez, laissez le reste vide, "
                      "puis importez le classeur entier dans Carène "
                      "(« Création du navire » → « Par où commencer »)."))
    lignes_lm.append(("Comment faire",
                      "Survolez un intitulé de colonne : un commentaire donne "
                      "l'unité, le repère, si le renseignement est obligatoire "
                      "et la valeur du navire d'exemple. N'ajoutez ni ne "
                      "renommez les colonnes sans nécessité ; une colonne "
                      "inconnue est simplement ignorée, et signalée."))
    lignes_lm.append(("Ordre conseillé",
                      " → ".join(n for n, _d in FEUILLES_DESCRIPTION)))
    lignes_lm.append(("", ""))
    lignes_lm.append(("LES FEUILLES", ""))
    lignes_lm.extend(FEUILLES_DESCRIPTION)
    lignes_lm.append(("", ""))
    lignes_lm.append(("CONVENTIONS ET REPÈRES",
                      conventions.get("a_quoi", "")))
    for nom, unite, _requis, exemple in conventions.get("format_rows", []):
        lignes_lm.append((nom, f"{unite} — exemple : {exemple}"))
    for note in conventions.get("notes", []):
        lignes_lm.append(("À retenir", note))
    for r, (a, b) in enumerate(lignes_lm, start=1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = gras
        ca.alignment = Alignment(vertical="top")
        cb = ws.cell(row=r, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)

    # ------------------------------------------------------- clé / valeur
    d = draft
    cle_valeur(F_IDENTIFICATION, CHAMPS_IDENTIFICATION + CHAMPS_DIMENSIONS,
               {**(d.identification if d else {}), **(d.dimensions if d else {})},
               "identification")
    cle_valeur(F_LEGE, CHAMPS_LEGE, d.lege if d else {}, "lege")
    cle_valeur(F_VENT, CHAMPS_VENT, d.vent if d else {}, "vent")

    # ------------------------------------------------------- tables
    table(F_PROFILS_VENT, SPEC_PROFILS_VENT, "vent",
          d.profils_vent_rows if d else None)
    couples = [{"n": n, "x_m": x} for n, x in (d.couples_paires() if d else [])]
    table(F_COUPLES, SPEC_COUPLES, "couples", couples)
    table(F_HYDRO, SPEC_HYDRO, "hydro", d.hydro_rows if d else None)

    angles = ([float(a) for a in d.kn_angles] if d and d.kn_angles
              else ANGLES_KN_TYPE)
    ws = feuille(F_KN)
    noms_kn = _entetes_kn(angles)
    com_kn = {
        "Assiette_m": "Unité / repère : m, positive sur l'arrière.\n"
                      "Renseignement : obligatoire.\nUne seule feuille pour "
                      "toutes les assiettes : c'est cette colonne qui les "
                      "distingue.\nExemple : -0.5, 0, 0.5, 1, 1.5",
        "Deplacement_t": "Unité / repère : t.\nRenseignement : obligatoire.\n"
                         "Exemple : 2600.0",
    }
    exemples_kn = {5: "0.475", 60: "5.047"}
    for a in angles:
        com_kn[f"KN_{float(a):g}"] = (
            f"Bras de levier de forme KN à {float(a):g}° de gîte.\n"
            "Unité / repère : m, compté depuis la quille (KG nul).\n"
            "Renseignement : obligatoire pour chaque angle tabulé ; une ligne "
            "dont une case d'angle est vide est rejetée.\nAjoutez ou retirez "
            "des colonnes selon votre table : l'angle est lu dans l'intitulé."
            + (f"\nExemple : {exemples_kn[int(a)]}"
               if int(a) in exemples_kn and float(a) == int(a) else ""))
    entetes(ws, noms_kn, com_kn)
    for r, ligne in enumerate(d.kn_rows if d else [], start=2):
        for i, n in enumerate(noms_kn, start=1):
            ws.cell(row=r, column=i, value=_valeur_cellule(ligne.get(n)))

    # la LISTE des capacités a ses propres colonnes (une ligne par capacité),
    # décrites à part de la table de jaugeage (une ligne par remplissage)
    table(F_CAPACITES, SPEC_CAPACITES, "capacites", d.capacites if d else None,
          cle_format="format_rows_liste")

    # une feuille de jaugeage par capacité quand le classeur est rempli, sinon
    # le seul modèle à dupliquer. Une capacité déclarée SANS jaugeage reçoit sa
    # feuille vide : le classeur montre ainsi ce qui reste à remplir, au lieu
    # de laisser croire que la capacité est complète.
    if d and (d.jauges or d.capacites):
        faites = set()
        for nom, lignes in d.jauges.items():
            table(nom_feuille_jauge(nom), SPEC_JAUGE, "capacites", lignes)
            faites.add(nom)
        for meta in d.capacites:
            nom = str(meta.get("Nom", "") or "")
            if nom and nom not in faites:
                table(nom_feuille_jauge(nom), SPEC_JAUGE, "capacites")
                faites.add(nom)
    else:
        table(F_JAUGE_EXEMPLE, SPEC_JAUGE, "capacites")

    table(F_ENVAHISSEMENT, SPEC_ENVAHISSEMENT, "envahissement",
          d.envahissement if d else None)
    table(F_CAS, SPEC_CAS_REFERENCE, "validation",
          d.cas_reference if d else None)

    dossier = os.path.dirname(os.path.abspath(chemin_xlsx))
    if dossier:
        os.makedirs(dossier, exist_ok=True)
    wb.save(chemin_xlsx)
    return wb.sheetnames


def nom_feuille_jauge(nom_capacite):
    """« WB 3B » → « Jauge_WB_3B » (31 caractères au plus : la limite d'Excel).

    Le souligné remplace l'espace ET la barre oblique, comme pour les fichiers
    de `jauges/` : c'est la même convention, pour que le nom de la feuille se
    retrouve dans le dossier écrit.
    """
    propre = str(nom_capacite).replace(" ", "_").replace("/", "_")
    return (PREFIXE_JAUGE + propre)[:31]


def capacite_de_feuille(nom_feuille, capacites_declarees=()):
    """Nom de capacité porté par une feuille « Jauge_… », ou None.

    Le nom écrit dans la feuille a perdu ses espaces et ses barres obliques :
    on le rapproche d'abord des capacités DÉCLARÉES (feuille « Capacites »),
    qui font foi ; à défaut, un souligné redevient un espace.
    """
    m = re.match(r"^\s*(?:jauge|jaugeage|sounding|ullage)\s*[_\- ]\s*(.+)$",
                 str(nom_feuille), re.IGNORECASE)
    if not m:
        return None
    brut = m.group(1).strip()
    cle = brut.replace(" ", "_").replace("/", "_")
    for nom in capacites_declarees:
        if str(nom).replace(" ", "_").replace("/", "_") == cle:
            return str(nom)
    return brut.replace("_", " ")


# ------------------------------------------------------------- gabarits CSV
def ecrire_modeles_csv(dossier, draft=None):
    """Écrit un CSV par feuille du classeur, plus un LISEZ-MOI.md.

    Même contenu, mêmes intitulés : c'est le chemin de secours pour un poste
    sans tableur complet. Les commentaires de cellule, eux, n'ont pas
    d'équivalent en CSV — ils sont repris en texte dans le LISEZ-MOI.
    """
    os.makedirs(dossier, exist_ok=True)
    etapes, conventions = _etapes()
    ecrits = []

    def ecrire(nom, colonnes, lignes=None):
        chemin = os.path.join(dossier, nom + ".csv")
        with open(chemin, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(colonnes)
            for ligne in lignes or []:
                w.writerow([ligne.get(c, "") for c in colonnes])
        ecrits.append(nom + ".csv")

    d = draft
    champs_id = CHAMPS_IDENTIFICATION + CHAMPS_DIMENSIONS
    valeurs_id = {**(d.identification if d else {}), **(d.dimensions if d else {})}
    for nom, champs, valeurs in (
            (F_IDENTIFICATION, champs_id, valeurs_id),
            (F_LEGE, CHAMPS_LEGE, d.lege if d else {}),
            (F_VENT, CHAMPS_VENT, d.vent if d else {})):
        lignes = [{"champ": cle, "valeur": (valeurs or {}).get(cle, ""),
                   "libelle": libelle + (" *" if requis else "")}
                  for cle, libelle, _u, requis in champs]
        ecrire(nom, ["champ", "valeur", "libelle"], lignes)

    ecrire(F_PROFILS_VENT, [c.name for c in SPEC_PROFILS_VENT],
           d.profils_vent_rows if d else None)
    ecrire(F_COUPLES, [c.name for c in SPEC_COUPLES],
           [{"n": n, "x_m": x} for n, x in (d.couples_paires() if d else [])])
    ecrire(F_HYDRO, [c.name for c in SPEC_HYDRO], d.hydro_rows if d else None)
    angles = ([float(a) for a in d.kn_angles] if d and d.kn_angles
              else ANGLES_KN_TYPE)
    ecrire(F_KN, _entetes_kn(angles), d.kn_rows if d else None)
    ecrire(F_CAPACITES, [c.name for c in SPEC_CAPACITES],
           d.capacites if d else None)
    if d and d.jauges:
        for nom, lignes in d.jauges.items():
            ecrire(nom_feuille_jauge(nom), [c.name for c in SPEC_JAUGE], lignes)
    else:
        ecrire(F_JAUGE_EXEMPLE, [c.name for c in SPEC_JAUGE])
    ecrire(F_ENVAHISSEMENT, [c.name for c in SPEC_ENVAHISSEMENT],
           d.envahissement if d else None)
    ecrire(F_CAS, [c.name for c in SPEC_CAS_REFERENCE],
           d.cas_reference if d else None)

    lignes = [
        "# Gabarits de reprise d'un dossier de stabilité",
        "",
        "Un fichier par table, mêmes intitulés de colonnes que le classeur "
        "Excel — c'est le chemin de secours pour un poste sans tableur. Les "
        "fichiers s'importent un par un dans la fenêtre « Création du "
        "navire », à l'étape correspondante.",
        "",
        "| Fichier | Contenu |",
        "| --- | --- |",
    ]
    lignes += [f"| `{nom}.csv` | {desc} |" for nom, desc in FEUILLES_DESCRIPTION]
    lignes += ["", "## Conventions et repères", "",
               conventions.get("a_quoi", ""), ""]
    lignes += [f"- **{nom}** — {unite}, {requis} (exemple : {ex})"
               for nom, unite, requis, ex in conventions.get("format_rows", [])]
    lignes += [""] + [f"> {n}" for n in conventions.get("notes", [])]
    lignes += ["", "## Colonnes attendues, table par table", ""]
    for titre, spec, cle, cle_format in (
            (F_HYDRO, SPEC_HYDRO, "hydro", "format_rows"),
            (F_CAPACITES, SPEC_CAPACITES, "capacites", "format_rows_liste"),
            (F_JAUGE_EXEMPLE, SPEC_JAUGE, "capacites", "format_rows"),
            (F_ENVAHISSEMENT, SPEC_ENVAHISSEMENT, "envahissement",
             "format_rows"),
            (F_COUPLES, SPEC_COUPLES, "couples", "format_rows"),
            (F_PROFILS_VENT, SPEC_PROFILS_VENT, "vent", "format_rows"),
            (F_CAS, SPEC_CAS_REFERENCE, "validation", "format_rows")):
        legendes = _apparier(spec, (etapes.get(cle) or {}).get(cle_format) or [])
        lignes += [f"### {titre}.csv", ""]
        for col in spec:
            leg = legendes.get(col.name)
            detail = (f"{leg[0]}, {leg[1]}" if leg else
                      ("obligatoire" if col.required else "facultatif"))
            exemple = f" — exemple : {leg[2]}" if leg and leg[2] not in ("", "—") else ""
            lignes.append(f"- `{col.name}` — {detail}{exemple}")
        lignes.append("")
    chemin = os.path.join(dossier, "LISEZ-MOI.md")
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n".join(lignes) + "\n")
    ecrits.append("LISEZ-MOI.md")
    return ecrits


# ------------------------------------------------------------- lecture
def _statut_import(rapport):
    """Résumé lisible d'un compte rendu d'import de table."""
    bouts = []
    for cle, libelle in (("lignes", "ligne(s)"), ("points", "point(s)"),
                         ("capacites", "capacité(s)"), ("couples", "couple(s)"),
                         ("cas", "cas")):
        if rapport.get(cle) is not None:
            bouts.append(f"{rapport[cle]} {libelle} lue(s)")
    if rapport.get("assiettes"):
        bouts.append("assiettes " + ", ".join(f"{t:+.3f}"
                                              for t in rapport["assiettes"]))
    if rapport.get("angles"):
        bouts.append("angles " + ", ".join(f"{a:g}°" for a in rapport["angles"]))
    if rapport.get("ignorees"):
        bouts.append("colonnes ignorées : " + ", ".join(
            str(c) for c in rapport["ignorees"][:6]))
    if rapport.get("rejetees"):
        bouts.append(f"{len(rapport['rejetees'])} ligne(s) rejetée(s)")
    return " ; ".join(bouts) or "feuille lue"


def importer_classeur(draft, chemin_xlsx):
    """Lit toutes les feuilles d'un classeur rempli dans `draft`.

    Rend [(feuille, statut, message)] avec statut dans « ok », « vide »,
    « ignoree », « erreur » — une feuille vide n'est pas une erreur, une
    feuille inconnue non plus : on le DIT, et on continue. Un classeur à demi
    rempli doit pouvoir être importé dix fois de suite pendant que l'officier
    le complète.
    """
    from .navire_draft import ouvrir_classeur

    wb = ouvrir_classeur(chemin_xlsx)
    try:
        feuilles = {ws.title: ws for ws in wb.worksheets}
        par_norme = {normalize(t): t for t in feuilles}
        vus = set()
        rapport = []

        def trouver(nom):
            titre = par_norme.get(normalize(nom))
            if titre is None:
                return None, None
            vus.add(titre)
            return titre, table_de_feuille(feuilles[titre])

        def poser(titre, statut, message):
            rapport.append((titre, statut, message))

        def lire_table(nom, charger):
            titre, tbl = trouver(nom)
            if titre is None:
                return
            headers, rows = tbl
            if not rows:
                poser(titre, "vide", "feuille non remplie — ignorée")
                return
            try:
                res = charger(headers, rows)
            except Exception as e:
                poser(titre, "erreur", str(e))
                return
            poser(titre, "ok", _statut_import(res or {}))

        def lire_cle_valeur(nom, champs, cible, quoi):
            titre, tbl = trouver(nom)
            if titre is None:
                return
            headers, rows = tbl
            connus = {c[0]: c for c in champs}
            par_norme_champ = {normalize(c[0]): c[0] for c in champs}
            par_norme_champ.update({normalize(c[1]): c[0] for c in champs})
            i_champ, i_valeur = 0, 1
            for i, h in enumerate(headers):
                hn = normalize(h)
                if hn in ("champ", "cle", "renseignement", "nom"):
                    i_champ = i
                elif hn in ("valeur", "value"):
                    i_valeur = i
            lus, inconnus = 0, []
            for r in rows:
                cle = str(r[i_champ]).strip() if i_champ < len(r) else ""
                brut = str(r[i_valeur]).strip() if i_valeur < len(r) else ""
                if not cle:
                    continue
                vrai = cle if cle in connus else par_norme_champ.get(normalize(cle))
                if vrai is None:
                    inconnus.append(cle)
                    continue
                if brut == "":
                    continue
                unite = connus[vrai][2]
                if unite == "texte":
                    cible[vrai] = brut
                else:
                    v = to_float(brut)
                    if v is None:
                        inconnus.append(f"{cle} (valeur illisible)")
                        continue
                    cible[vrai] = v
                lus += 1
            if not lus and not inconnus:
                poser(titre, "vide", "feuille non remplie — ignorée")
                return
            msg = f"{lus} {quoi} renseigné(s)"
            if inconnus:
                msg += " ; ignoré(s) : " + ", ".join(inconnus[:6])
            poser(titre, "ok" if lus else "vide", msg)

        # l'ordre est imposé ici, pas par le classeur : les capacités doivent
        # être connues avant leurs jaugeages (c'est la liste déclarée qui donne
        # le nom exact d'une feuille « Jauge_… »)
        lire_cle_valeur(F_IDENTIFICATION,
                        CHAMPS_IDENTIFICATION + CHAMPS_DIMENSIONS,
                        _cible_identification(draft), "renseignement(s)")
        lire_cle_valeur(F_LEGE, CHAMPS_LEGE, draft.lege, "valeur(s)")
        lire_cle_valeur(F_VENT, CHAMPS_VENT, draft.vent, "valeur(s)")
        lire_table(F_PROFILS_VENT, draft.charger_profils_vent)
        lire_table(F_COUPLES, draft.charger_couples)
        lire_table(F_HYDRO, draft.charger_hydro)
        lire_table(F_KN, draft.charger_kn)
        lire_table(F_CAPACITES, draft.charger_capacites)

        declarees = [c.get("Nom") for c in draft.capacites]
        for titre in list(feuilles):
            if titre in vus:
                continue
            nom_cap = capacite_de_feuille(titre, declarees)
            if nom_cap is None:
                continue
            vus.add(titre)
            headers, rows = table_de_feuille(feuilles[titre])
            if not rows:
                poser(titre, "vide", "feuille non remplie — ignorée")
                continue
            try:
                res = draft.charger_jauge(headers, rows, nom_cap)
            except Exception as e:
                poser(titre, "erreur", str(e))
                continue
            poser(titre, "ok", f"« {nom_cap} » : " + _statut_import(res or {}))

        lire_table(F_ENVAHISSEMENT, draft.charger_envahissement)
        lire_table(F_CAS, draft.charger_cas_reference)

        for titre in feuilles:
            if titre in vus or normalize(titre) == normalize(F_LISEZMOI):
                continue
            poser(titre, "ignoree", "feuille inconnue — rien n'en a été lu")

        draft.classeur_importe = os.path.basename(chemin_xlsx)
        return rapport
    finally:
        wb.close()


def _cible_identification(draft):
    """Un dict qui répartit les clés entre `identification` et `dimensions`.

    Les deux formulaires de la fenêtre sont distincts (et le manifeste les
    range séparément), mais l'utilisateur, lui, n'a qu'une feuille : « nom du
    navire » et « largeur » s'y suivent. Cette vue d'écriture fait le tri.
    """
    cles_dim = {c[0] for c in CHAMPS_DIMENSIONS}

    class _Repartiteur(dict):
        def __setitem__(self, cle, valeur):
            if cle in cles_dim:
                draft.dimensions[cle] = valeur
            else:
                draft.identification[cle] = valeur
            dict.__setitem__(self, cle, valeur)

    return _Repartiteur()
