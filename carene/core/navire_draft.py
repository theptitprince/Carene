# -*- coding: utf-8 -*-
"""Construction d'un dossier « navire virtuel » depuis des tables importées.

Ce module est le pendant *écriture* de `navire.py` (qui, lui, lit un dossier
déjà constitué). Il sert à la fenêtre « Création du navire » :

- lire un tableau CSV/Excel quelconque (`read_table`) ;
- reconnaître ses colonnes malgré des intitulés variables (`match_columns`) ;
- accumuler les données dans un `NavireDraft` ;
- calculer l'état d'avancement (`NavireDraft.status`) ;
- écrire le dossier navire (`NavireDraft.save`).

Aucune donnée propre à un navire n'est codée ici : seuls figurent le FORMAT du
dossier, les synonymes de colonnes, et les seuils RÉGLEMENTAIRES par défaut
(constantes universelles IMO/BV, pas des valeurs de navire) — que l'utilisateur
peut de toute façon modifier et qui sont écrits dans navire.json.
"""
from __future__ import annotations
from ..ecriture import ecriture_atomique

import csv
import json
import os
import re
import unicodedata

# --------------------------------------------------------------- utilitaires


def normalize(s: str) -> str:
    """Intitulé réduit à sa forme comparable : sans accent, sans ponctuation."""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def nom_de_dossier(nom: str) -> str:
    """Nom de navire → nom de DOSSIER transportable : « Sœur Océane » →
    « SOEUR_OCEANE ».

    Un dossier de navire se copie sur une clé et s'ouvre sur un autre poste :
    accents, espaces et ponctuation y sont autant d'occasions qu'un chemin ne
    se retrouve pas (encodage Windows, archive zip, ligne de commande). Les
    majuscules sont la convention du dossier livré (`navires/<NOM>`).
    """
    s = str(nom or "")
    for lig, remp in (("œ", "oe"), ("Œ", "OE"), ("æ", "ae"), ("Æ", "AE")):
        s = s.replace(lig, remp)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_").upper()
    return s or "NAVIRE"


def ouvrir_classeur(path):
    """Classeur Excel ouvert en lecture (valeurs, pas formules).

    Sorti de `read_table` pour que l'import d'un CLASSEUR entier (une feuille
    par table, voir `carene.core.classeur_navire`) n'ait pas à rouvrir le
    fichier une fois par feuille.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise RuntimeError(
            "La lecture des fichiers Excel demande le module openpyxl "
            "(pip install openpyxl). Vous pouvez aussi exporter la feuille "
            "en CSV.") from e
    return load_workbook(path, read_only=True, data_only=True)


def lignes_de_feuille(ws):
    """Lignes non vides d'une feuille déjà ouverte, en texte.

    `read_table` ne sait lire que LA PREMIÈRE feuille d'un classeur : depuis
    que le navire tient dans un seul classeur, chaque table est une feuille
    qu'on va chercher par son nom. Cette fonction est la brique commune.
    """
    rows = [["" if c is None else str(c) for c in r]
            for r in ws.iter_rows(values_only=True)]
    return [r for r in rows if any(str(c).strip() for c in r)]


def table_de_feuille(ws):
    """(intitulés, lignes) d'une feuille de classeur.

    Mêmes conventions que `read_table`, mais une feuille SANS DONNÉE (rien, ou
    seulement ses intitulés) rend ([], []) ou (intitulés, []) au lieu de lever
    une erreur : dans un classeur type, les feuilles non remplies sont la
    règle, pas l'exception.
    """
    rows = lignes_de_feuille(ws)
    if not rows:
        return [], []
    return [str(c).strip() for c in rows[0]], rows[1:]


def read_table(path):
    """(intitulés, lignes) d'un CSV (délimiteur deviné) ou d'un XLSX.

    Les lignes sont des listes de chaînes ; la conversion en nombre est faite
    plus tard, colonne par colonne, pour pouvoir signaler proprement une
    cellule non numérique.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        wb = ouvrir_classeur(path)
        rows = lignes_de_feuille(wb[wb.sheetnames[0]])
        wb.close()
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(8192)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = [list(r) for r in csv.reader(f, dialect)]
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("Fichier vide.")
    return [str(c).strip() for c in rows[0]], rows[1:]


def to_float(text):
    """Nombre lu de façon tolérante : virgule décimale, espaces, signe unicode."""
    if text is None:
        return None
    s = str(text).strip().replace("−", "-").replace(" ", "")
    s = s.replace(" ", "")
    if not s:
        return None
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    elif "," in s and "." in s:
        # les deux : le DERNIER est le séparateur décimal — « 1.129,24 »
        # (européen) comme « 1,129.24 » (anglais). Retirer toujours la
        # virgule lisait le premier 1,12924 (D-82).
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    try:
        x = float(s)
    except ValueError:
        return None
    # « nan », « inf » : des mots, pas des valeurs de table
    return x if (x == x and x not in (float("inf"), float("-inf"))) else None


class Col:
    """Une colonne attendue : nom canonique, synonymes acceptés, obligatoire."""

    def __init__(self, name, aliases, required=False):
        self.name = name
        self.aliases = [normalize(a) for a in ([name] + list(aliases))]
        self.required = required

    def matches(self, header_norm):
        if header_norm in self.aliases:
            return 2                      # correspondance exacte
        for a in self.aliases:
            if len(a) >= 3 and (header_norm.startswith(a) or a in header_norm):
                return 1                  # correspondance partielle
        return 0


# --------------------------------------------------------------- spécifications
SPEC_HYDRO = [
    Col("Assiette_m", ["assiette", "trim"]),
    Col("TE_milieu_m", ["te milieu", "tirant d'eau moyen", "tirant d'eau",
                        "draft", "mean draft", "tmoy", "tm"], required=True),
    Col("Deplacement_t", ["deplacement", "displacement", "depl", "wght",
                          "weight", "poids"], required=True),
    Col("LCB_m", ["lcb", "centre de carene", "long center of buoyancy"],
        required=True),
    Col("KMt_m", ["kmt", "km transversal", "km", "kmtransversal"], required=True),
    Col("LCF_m", ["lcf", "centre de flottaison", "long center of flotation"]),
    Col("VCB_m", ["vcb", "kb", "centre de carene vertical"]),
    Col("KMl_m", ["kml", "km longitudinal"]),
    Col("MCT_tm_cm", ["mct", "moment unitaire", "moment to trim"]),
    Col("TPC_t_cm", ["tpc", "tonnes par cm", "tonnes per cm"]),
    Col("WPA_m2", ["wpa", "aire de flottaison", "waterplane area"]),
    Col("Volume_m3", ["volume", "vol"]),
    Col("TE_AR_m", ["te ar", "tirant d'eau arriere", "draft aft", "ta"]),
    Col("TE_AV_m", ["te av", "tirant d'eau avant", "draft fwd", "tf"]),
    Col("BMT_m", ["bmt", "bm transversal"]),
    Col("BML_m", ["bml", "bm longitudinal"]),
    Col("WSA_m2", ["wsa", "surface mouillee", "wetted surface"]),
]

SPEC_JAUGE = [
    Col("Remplissage_pc", ["remplissage", "filling", "taux", "pc", "percent"],
        required=True),
    Col("Volume_m3", ["volume", "vol"], required=True),
    Col("Poids_t", ["poids", "weight", "wght", "masse"], required=True),
    Col("LCG_m", ["lcg", "centre de gravite longitudinal"], required=True),
    Col("TCG_m", ["tcg", "centre de gravite transversal"], required=True),
    Col("VCG_m", ["vcg", "centre de gravite vertical", "kg"], required=True),
    Col("FSM_tm", ["fsm", "moment de carene liquide", "free surface moment",
                   "fsmt"], required=True),
    Col("Sondage_m", ["sondage", "sounding"]),
    Col("Creux_m", ["creux", "ullage"]),
    Col("VCG_corrige_m", ["vcg corrige", "vcgt", "vcg corr"]),
]

SPEC_CAPACITES = [
    Col("Nom", ["nom", "name", "code", "capacite", "tank"], required=True),
    Col("Type", ["type", "contenu", "content", "liquide"]),
    Col("Groupe", ["groupe", "group", "weight group", "famille"]),
    Col("Densite", ["densite", "density", "masse volumique"]),
    Col("Volume_net_m3", ["volume net", "volume", "capacite m3"]),
    Col("Perm_pc", ["permeabilite", "perm"]),
    Col("Poids_t", ["poids", "weight"]),
    Col("LCG_m", ["lcg"]),
    Col("TCG_m", ["tcg"]),
    Col("VCG_m", ["vcg"]),
    Col("FSM_max_tm", ["fsm max", "fsm", "moment de carene liquide max"]),
    Col("Baffles", ["baffles", "cloisons"]),
    Col("Ballast", ["ballast", "ballastable", "wb"]),
]

SPEC_ENVAHISSEMENT = [
    Col("Repere", ["repere", "nom", "code", "id", "opening"], required=True),
    Col("Fonction", ["fonction", "description", "designation", "type"]),
    Col("X_m", ["x", "abscisse", "lcg", "position x"], required=True),
    Col("Y_m", ["y", "ordonnee", "tcg", "position y"], required=True),
    Col("Z_m", ["z", "hauteur", "vcg", "position z"], required=True),
]

# Surfaces exposées au vent, configuration de voilure par configuration et
# tirant d'eau par tirant d'eau. C'est ce qui manquait pour évaluer les
# critères de vent à l'écran : l'aire au vent n'est ni une constante du navire
# ni une formule, elle dépend de la toile dehors ET du tirant d'eau. Une ligne
# par mesure du dossier ; `carene.core.voilure` interpole entre deux lignes
# d'une même configuration et refuse de prolonger au-delà.
SPEC_PROFILS_VENT = [
    Col("Profil", ["profil", "voilure", "configuration", "config", "sail",
                   "sails", "id"], required=True),
    Col("Source", ["source", "cas", "origine", "reference", "provenance"]),
    Col("TE_milieu_m", ["te milieu", "tirant d'eau moyen", "tirant d'eau",
                        "hmp", "draft", "mean draft"], required=True),
    Col("Windage_area_m2", ["windage area", "aire au vent", "aire windage",
                            "surface au vent", "aire exposee au vent"],
        required=True),
    Col("Windage_V_m", ["windage v", "vcw", "hauteur du centre au vent",
                        "centre au vent", "windage vcg"], required=True),
    Col("Lateral_plane_area_m2", ["lateral plane area", "aire plan de derive",
                                  "plan de derive", "aire derive"]),
    Col("Lateral_plane_V_m", ["lateral plane v", "vclp",
                              "centre du plan de derive", "derive v"],
        required=True),
    Col("Z_windage_lateral_m", ["z windage lateral", "z", "bras de vent",
                                "bras windage", "levier de vent"]),
]

# Table des couples : c'est elle qui permet de caler un plan « sur le couple
# 35 » plutôt qu'en mesurant à l'écran (D-11, `app_paths.table_couples`).
SPEC_COUPLES = [
    Col("n", ["n", "no", "num", "numero", "couple", "frame", "membrure"],
        required=True),
    Col("x_m", ["x", "abscisse", "position", "distance", "x depuis c0"],
        required=True),
]

# Cas de référence : un cas de chargement du dossier approuvé, recopié avec
# ses résultats imprimés, pour que le logiciel les recalcule et compare.
# `VCG_solide_m` OU `VCG_corrige_m` suffit (l'un se déduit de l'autre par le
# FSM) : la vérification de ce couple-là se fait à l'import, pas ici.
SPEC_CAS_REFERENCE = [
    Col("Cas", ["cas", "code", "condition", "load case", "loadcase"],
        required=True),
    Col("Titre", ["titre", "libelle", "designation", "title", "description"]),
    Col("Poids_total_t", ["poids total", "wght total", "deplacement total",
                          "displacement", "deplacement", "poids", "masse"],
        required=True),
    Col("LCG_m", ["lcg", "lcg total", "centre de gravite longitudinal"],
        required=True),
    Col("TCG_m", ["tcg", "tcg total", "centre de gravite transversal"]),
    Col("VCG_solide_m", ["vcg solide", "kg solide", "vcg", "kg",
                         "centre de gravite vertical"]),
    Col("VCG_corrige_m", ["vcg corrige", "kg corrige", "vcgc", "vcg total"]),
    Col("FSM_tm", ["fsm", "fsmt", "fsm total", "moment de carene liquide"],
        required=True),
    Col("TE_milieu_m", ["te milieu", "tirant d'eau moyen", "tirant d'eau",
                        "hmp", "draft"]),
    Col("Assiette_m", ["assiette", "trim"]),
    Col("GM_corrige_m", ["gm corrige", "gmtc", "gm"]),
    Col("GZmax_m", ["gzmax", "gz max", "max gz"]),
    Col("Angle_GZmax_deg", ["angle gzmax", "angle du gzmax", "angle gz max",
                            "angle max gz"]),
]

# Tolérances de comparaison d'un cas de référence — celles que
# `tools/valider_reference.py` applique aux 15 cas du dossier de référence. Elles
# valent pour tout dossier : ce sont les ordres de grandeur en deçà desquels
# un écart s'explique par l'interpolation des tables et l'arrondi d'impression,
# pas par une colonne mal reprise. La constante est ici (et non dans l'outil)
# pour que la fenêtre « Création du navire » et le rejeu en ligne de commande
# ne puissent pas diverger.
TOLERANCES_CAS = {
    "draft_m": 0.02,
    "trim_m": 0.06,
    "gm_m": 0.015,
    "gzmax_m": 0.06,
    "angle_gzmax_deg": 3.0,
    "aire_m_rad": 0.010,
}

# Les grandeurs comparées : (colonne attendue, libellé, clé de tolérance).
# Une colonne absente du fichier n'est simplement pas comparée — on ne compare
# que ce que le dossier a imprimé.
GRANDEURS_CAS = [
    ("TE_milieu_m", "Tirant d'eau moyen (m)", "draft_m"),
    ("Assiette_m", "Assiette (m)", "trim_m"),
    ("GM_corrige_m", "GM corrigé (m)", "gm_m"),
    ("GZmax_m", "GZmax (m)", "gzmax_m"),
    ("Angle_GZmax_deg", "Angle du GZmax (°)", "angle_gzmax_deg"),
]

# seuils RÉGLEMENTAIRES par défaut (constantes IMO/BV, pas des données navire)
CRITERES_DEFAUT = {
    "is2008_general": {
        "reference": "IMO IS Code 2008, Part A, §2.2",
        "aire_0_30_m_rad_min": 0.055,
        "aire_0_40_m_rad_min": 0.09,
        "aire_30_40_m_rad_min": 0.03,
        "gz_max_m_min": 0.2,
        "angle_gz_max_deg_min": 25.0,
        "gm0_corrige_m_min": 0.15,
    },
    "nr500_voile": {
        "reference": "BV NR500 (navires à voile)",
        "gm0_corrige_m_min": 0.3,
    },
    "nr500_voile_sous_voile": {
        "reference": "BV NR500 — critère de gîte sous voile",
        "gm0_corrige_m_min": 0.3,
        "gz_a_50deg_ou_gzmax_m_min": 0.5,
        "angle_gzmax_deg_min": 50.0,
        # aire entre la courbe GZ et le bras de gîte sous vent, de l'angle
        # statique à l'angle d'envahissement (cf. criteria.evaluate_nr500_voilier)
        "aire_gz_vent_m_rad_min": 0.065,
        "angle_statique_vent_deg_max": 20.0,
        "fraction_angle_immersion_pont": 0.90,
        # pression de vent = k · V² : k = ½·ρ_air avec ρ_air = 1,222 kg/m³ —
        # la convention lue dans le recueil du navire de référence (P/V² = 0,611 sur ses
        # six cas sous voile). Ne sert qu'à afficher la vitesse (D-58). Le
        # coefficient de forme F = C · P · A (1,10 sur le navire de référence) n'a PAS de
        # défaut : il dépend du logiciel du dossier.
        "pression_par_v2_pa_s2_m2": 0.611,
    },
    "critere_meteo": {
        "reference": "IMO IS Code 2008, Part A, §2.3 (severe wind and rolling)",
        "pression_vent_pa": 504.0,
        "angle_gite_vent_stable_deg_max": 16.0,
    },
}

CHAMPS_IDENTIFICATION = [
    ("nom", "Nom du navire", "texte", True),
    ("imo", "Numéro OMI / immatriculation", "texte", False),
    ("type", "Type de navire", "texte", False),
    ("pavillon", "Pavillon", "texte", False),
    # Convention de carène liquide du DOSSIER, pas du logiciel : « reel » = FSM
    # interpolé au remplissage courant (recueil approuvé du navire de référence,
    # LOCOPIAS),
    # « max » = FSM maximal de chaque capacité, plus conservatif. Écrite en
    # tête de navire.json (`convention_fsm`), là où le moteur la lit ; sans
    # elle un navire rebâti retombait sur « max » sans qu'on l'ait choisi.
    ("convention_fsm", "Carène liquide : « reel » (FSM au remplissage) "
     "ou « max » (FSM maximal)", "texte", False),
]
CONVENTIONS_FSM = ("reel", "max")

CHAMPS_DIMENSIONS = [
    ("longueur_hors_tout_m", "Longueur hors tout", "m", False),
    ("longueur_entre_pp_hydro_m", "Longueur entre perpendiculaires", "m", True),
    ("largeur_hors_membres_m", "Largeur hors membres (B)", "m", True),
    ("creux_sur_quille_m", "Creux sur quille (D)", "m", True),
    ("aire_quilles_anti_roulis_m2", "Aire des quilles anti-roulis", "m²", False),
    # abscisses des perpendiculaires dans le repère du dossier : sans elles,
    # le tirant d'eau local (donc l'angle d'envahissement) ignore l'assiette
    ("x_perpendiculaire_ar_m", "X de la perpendiculaire arrière", "m", False),
    ("x_perpendiculaire_av_m", "X de la perpendiculaire avant", "m", False),
    # la densité de l'eau dans laquelle les tables ont été calculées : elle ne
    # sert qu'au relevé de tirants d'eau (D-60), qui corrige de la densité du
    # port. Absente, 1,025 (eau de mer) est retenue et le dit.
    ("densite_eau_tables", "Densité de l'eau des tables hydrostatiques", "t/m³", False),
    # la marque d'été du certificat de franc-bord (D-78) : sans elle, la ligne
    # de charge n'est pas contrôlée — et Carène le dit
    ("tirant_eau_ete_m", "Tirant d'eau d'été (certificat de franc-bord)", "m", False),
]

CHAMPS_LEGE = [
    ("masse_t", "Masse lège", "t", True),
    ("lcg_m", "LCG (depuis l'origine longitudinale)", "m", True),
    ("vcg_m", "VCG (au-dessus de la quille)", "m", True),
    # facultatif : ≈ 0 pour un navire symétrique, souvent omis des dossiers
    ("tcg_m", "TCG (positif bâbord)", "m", False),
    # D'où viennent ces chiffres — en texte, facultatif. Sans ces trois
    # lignes, un navire rebâti depuis le classeur calcule pareil mais son
    # manifeste ne dit plus si c'est la masse mesurée ou celle du dossier
    # (D-16 : le navire de référence retient la masse de l'expérience et les centres du
    # dossier approuvé). Elles passent telles quelles dans navire.json.
    ("date_experience", "Date de l'expérience de stabilité", "texte", False),
    ("base", "Base retenue (masse mesurée, centres du dossier…)", "texte", False),
    ("source", "Source : références du dossier, notes", "texte", False),
]

# Surface au vent : PAS une propriété du navire — elle dépend du gréement (ou
# des superstructures) ET du tirant d'eau, donc de la condition de chargement.
# Ces champs sont donc FACULTATIFS : ils décrivent une configuration de
# référence, faute de mieux, et le critère météo n'est évalué que s'ils sont là.
CHAMPS_VENT = [
    ("aire_windage_m2", "Aire de la surface exposée au vent", "m²", False),
    ("vcw_m", "Hauteur du centre de la surface exposée", "m", False),
    ("aire_plan_derive_m2", "Aire du plan de dérive immergé", "m²", False),
    ("vclp_m", "Hauteur du centre du plan de dérive", "m", False),
]


def match_columns(headers, spec):
    """Associe les intitulés d'un fichier aux colonnes attendues.

    Retourne (mapping {nom_canonique: index}, ignorées, manquantes_obligatoires).
    Chaque colonne du fichier n'est utilisée qu'une fois, la meilleure
    correspondance l'emportant.
    """
    norms = [normalize(h) for h in headers]
    scored = []
    for col in spec:
        for i, hn in enumerate(norms):
            s = col.matches(hn)
            if s:
                scored.append((s, col.name, i))
    scored.sort(key=lambda t: -t[0])
    mapping, used_idx, used_col = {}, set(), set()
    for _, name, i in scored:
        if name in used_col or i in used_idx:
            continue
        mapping[name] = i
        used_col.add(name)
        used_idx.add(i)
    ignored = [headers[i] for i in range(len(headers)) if i not in used_idx]
    missing = [c.name for c in spec if c.required and c.name not in mapping]
    return mapping, ignored, missing


def rows_to_dicts(rows, mapping, numeric_except=(), spec=None, rejetees=None):
    """Convertit les lignes brutes en dictionnaires selon `mapping`.

    Les colonnes citées dans `numeric_except` restent du texte ; les autres
    sont converties en nombre. Une ligne n'est ignorée que si une colonne
    OBLIGATOIRE (`required` dans `spec`) est vide ou illisible : une cellule
    vide dans une colonne facultative (LCF, TPC, sondage…) laissait auparavant
    tomber toute la ligne en silence — la clé est simplement absente du dict.
    Sans `spec`, toutes les colonnes du mapping sont tenues pour obligatoires
    (comportement historique). Si `rejetees` est une liste, les numéros de
    lignes ignorées (1 = première ligne de données) y sont ajoutés.
    """
    if spec is None:
        obligatoires = set(mapping)
    else:
        obligatoires = {c.name for c in spec if c.required}
    out = []
    for num, r in enumerate(rows, start=1):
        d = {}
        ok = True
        for name, idx in mapping.items():
            raw = r[idx] if idx < len(r) else ""
            if name in numeric_except:
                d[name] = str(raw).strip()
                continue
            v = to_float(raw)
            if v is None:
                if name in obligatoires:
                    ok = False
                    break
                continue
            d[name] = v
        if ok and d:
            out.append(d)
        elif rejetees is not None:
            rejetees.append(num)
    return out


# --------------------------------------------------------------- brouillon
class NavireDraft:
    """Un navire en cours de constitution, en mémoire."""

    def __init__(self):
        self.identification = {}
        self.dimensions = {}
        self.reperes = {
            "origine_longitudinale": "perpendiculaire arrière (PPAR)",
            "x_positif": "vers l'avant",
            "y_positif": "vers bâbord",
            "z_positif": "vers le haut, origine à la ligne de base (quille)",
        }
        self.lege = {}
        self.criteres = json.loads(json.dumps(CRITERES_DEFAUT))
        # Les critères VOILE (NR500) sont opt-in : actifs par défaut, ils
        # imposeraient GM ≥ 0,30 m à un cargo à moteur et le déclareraient
        # NON CONFORME à tort. Un navire neuf reçoit IS2008 + météo ; les
        # NR500 se cochent à l'étape « Critères » si le navire porte des voiles.
        self.criteres_actifs = {k: not k.startswith("nr500")
                                for k in CRITERES_DEFAUT}
        # LE PROFIL ET LES RÉGLEMENTATIONS (3.0, D-79). Le profil (type de
        # navire, propulsion, vitesse de service, cargaisons particulières)
        # fait PROPOSER des réglementations de la bibliothèque ; on retient
        # celles qu'on coche. `reglements_actifs` à None : rien n'a encore été
        # choisi dans la bibliothèque, on suit les anciens jeux de critères
        # (`criteres_actifs`) — c'est le cas de tout dossier d'avant la 3.0.
        self.profil = {}
        self.reglements_actifs = None
        self.hydro_rows = []          # dicts, avec Assiette_m
        self.kn_rows = []             # dicts, avec Assiette_m, Deplacement_t, KN_*
        self.kn_angles = []
        self.capacites = []           # dicts (meta)
        self.jauges = {}              # nom -> [dicts]
        self.envahissement = []       # dicts
        self.vent = {}
        # table des surfaces au vent par configuration de voilure et par
        # tirant d'eau (profils_vent.csv) — voir SPEC_PROFILS_VENT. Les quatre
        # champs `self.vent` de l'étape « Surface exposée au vent » décrivent
        # UNE configuration de référence, faute de mieux ; cette table-ci les
        # décrit toutes, et c'est elle que lisent les critères de vent.
        self.profils_vent_rows = []   # dicts (SPEC_PROFILS_VENT)
        # table des couples : x par n (index = numéro de couple), nan pour un
        # couple absent — même forme que `app_paths.table_couples`
        self.couples = []
        self.cas_reference = []       # dicts (SPEC_CAS_REFERENCE)
        self.source_path = ""
        self._extra_manifest = {}
        # mémoire de travail de la fenêtre, jamais enregistrée : le classeur
        # qui a servi au dernier import, et le dernier rejeu des cas
        self.classeur_importe = ""
        self.rejeu_cas = None

    # ------------------------------------------------------------ imports
    def import_hydro(self, path, assiette=None):
        """Importe une table hydrostatique. `assiette` s'applique si le fichier
        ne porte pas de colonne d'assiette. Retourne un compte rendu."""
        headers, rows = read_table(path)
        return self.charger_hydro(headers, rows, assiette)

    def charger_hydro(self, headers, rows, assiette=None):
        """Même chose depuis une table déjà lue (feuille de classeur)."""
        mapping, ignored, missing = match_columns(headers, SPEC_HYDRO)
        if missing:
            raise ValueError("Colonnes obligatoires introuvables : "
                             + ", ".join(missing))
        has_trim = "Assiette_m" in mapping
        if not has_trim and assiette is None:
            raise ValueError(
                "Ce fichier ne contient pas de colonne d'assiette : précisez à "
                "quelle assiette il correspond.")
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, spec=SPEC_HYDRO, rejetees=rejetees)
        if not has_trim:
            for d in dicts:
                d["Assiette_m"] = float(assiette)
        else:
            # l'assiette n'est « obligatoire » que si le fichier en a une
            # colonne : une ligne qui ne la renseigne pas est inexploitable
            sans = [d for d in dicts if "Assiette_m" not in d]
            if sans:
                if assiette is not None:
                    for d in sans:
                        d["Assiette_m"] = float(assiette)
                else:
                    dicts = [d for d in dicts if "Assiette_m" in d]
                    rejetees.append(f"{len(sans)} sans assiette")
        trims = sorted({round(d["Assiette_m"], 4) for d in dicts})
        # remplace les lignes de mêmes assiettes (ré-import d'un fichier corrigé)
        self.hydro_rows = [r for r in self.hydro_rows
                           if round(r["Assiette_m"], 4) not in trims]
        self.hydro_rows.extend(dicts)
        self.hydro_rows.sort(key=lambda r: (r["Assiette_m"], r["TE_milieu_m"]))
        return {"lignes": len(dicts), "assiettes": trims,
                "colonnes_reconnues": sorted(mapping), "ignorees": ignored,
                "rejetees": rejetees}

    def import_kn(self, path, assiette=None):
        """Importe des pantocarènes : une colonne de déplacement puis une
        colonne par angle de gîte (intitulé = l'angle)."""
        headers, rows = read_table(path)
        return self.charger_kn(headers, rows, assiette)

    def charger_kn(self, headers, rows, assiette=None):
        """Même chose depuis une table déjà lue (feuille de classeur)."""
        idx_depl = idx_trim = None
        angles = []                    # (angle, index)
        for i, h in enumerate(headers):
            hn = normalize(h)
            if idx_depl is None and any(
                    a in hn for a in ("deplacement", "displacement", "depl", "wght")):
                idx_depl = i
                continue
            if idx_trim is None and any(a in hn for a in ("assiette", "trim")):
                idx_trim = i
                continue
            m = re.search(r"(\d+(?:[.,]\d+)?)", h)
            if m and ("kn" in hn or "deg" in hn or hn.replace(".", "").isdigit()
                      or "gite" in hn or "heel" in hn):
                angles.append((float(m.group(1).replace(",", ".")), i))
        if idx_depl is None:
            raise ValueError("Colonne de déplacement introuvable.")
        if not angles:
            raise ValueError(
                "Aucune colonne d'angle de gîte reconnue : les intitulés "
                "doivent contenir l'angle (par exemple « KN 30 » ou « 30° »).")
        if idx_trim is None and assiette is None:
            raise ValueError(
                "Ce fichier ne contient pas de colonne d'assiette : précisez à "
                "quelle assiette il correspond.")
        angles.sort()
        dicts = []
        for r in rows:
            depl = to_float(r[idx_depl]) if idx_depl < len(r) else None
            if depl is None:
                continue
            trim = (to_float(r[idx_trim]) if idx_trim is not None
                    and idx_trim < len(r) else assiette)
            if trim is None:
                continue
            d = {"Assiette_m": float(trim), "Deplacement_t": depl}
            ok = True
            for ang, i in angles:
                v = to_float(r[i]) if i < len(r) else None
                if v is None:
                    ok = False
                    break
                # angle conservé EXACT : int() tronquait 7.5° en 7° et
                # décalait silencieusement toute la courbe GZ
                d[f"KN_{float(ang):g}"] = v
            if ok:
                dicts.append(d)
        if not dicts:
            raise ValueError("Aucune ligne exploitable dans ce fichier.")
        trims = sorted({round(d["Assiette_m"], 4) for d in dicts})
        self.kn_rows = [r for r in self.kn_rows
                        if round(r["Assiette_m"], 4) not in trims]
        self.kn_rows.extend(dicts)
        self.kn_rows.sort(key=lambda r: (r["Assiette_m"], r["Deplacement_t"]))
        self.kn_angles = sorted({float(a) for a, _ in angles}
                                | {float(a) for a in self.kn_angles})
        return {"lignes": len(dicts), "assiettes": trims,
                "angles": [float(a) for a, _ in angles]}

    def import_capacites(self, path):
        headers, rows = read_table(path)
        return self.charger_capacites(headers, rows)

    def charger_capacites(self, headers, rows):
        mapping, ignored, missing = match_columns(headers, SPEC_CAPACITES)
        if missing:
            raise ValueError("Colonnes obligatoires introuvables : "
                             + ", ".join(missing))
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, numeric_except=("Nom", "Type", "Groupe"),
                              spec=SPEC_CAPACITES, rejetees=rejetees)
        # une capacité sans nom n'est pas rattachable à une table de jaugeage
        dicts = [d for d in dicts if d.get("Nom")]
        if not dicts:
            raise ValueError("Aucune ligne exploitable.")
        self.capacites = dicts
        return {"capacites": len(dicts), "ignorees": ignored,
                "rejetees": rejetees}

    def import_jauge(self, path, nom=None):
        headers, rows = read_table(path)
        nom = nom or os.path.splitext(os.path.basename(path))[0].replace("_", " ")
        return self.charger_jauge(headers, rows, nom)

    def charger_jauge(self, headers, rows, nom):
        mapping, ignored, missing = match_columns(headers, SPEC_JAUGE)
        if missing:
            raise ValueError("Colonnes obligatoires introuvables : "
                             + ", ".join(missing))
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, spec=SPEC_JAUGE, rejetees=rejetees)
        if not dicts:
            raise ValueError("Aucune ligne exploitable.")
        dicts.sort(key=lambda r: r["Remplissage_pc"])
        self.jauges[nom] = dicts
        return {"capacite": nom, "lignes": len(dicts),
                "remplissage": (dicts[0]["Remplissage_pc"],
                                dicts[-1]["Remplissage_pc"]),
                "rejetees": rejetees}

    def import_envahissement(self, path):
        headers, rows = read_table(path)
        return self.charger_envahissement(headers, rows)

    def charger_envahissement(self, headers, rows):
        mapping, ignored, missing = match_columns(headers, SPEC_ENVAHISSEMENT)
        if missing:
            raise ValueError("Colonnes obligatoires introuvables : "
                             + ", ".join(missing))
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, numeric_except=("Repere", "Fonction"),
                              spec=SPEC_ENVAHISSEMENT, rejetees=rejetees)
        self.envahissement = dicts
        return {"points": len(dicts), "ignorees": ignored, "rejetees": rejetees}

    # ------------------------------------------------- surfaces au vent
    def import_profils_vent(self, path):
        """Importe la table des surfaces au vent (une ligne par configuration
        de voilure et par tirant d'eau)."""
        headers, rows = read_table(path)
        return self.charger_profils_vent(headers, rows)

    def charger_profils_vent(self, headers, rows):
        mapping, ignored, missing = match_columns(headers, SPEC_PROFILS_VENT)
        if missing:
            raise ValueError(
                "Colonnes obligatoires introuvables : " + ", ".join(missing)
                + " (attendu : la configuration de voilure, le tirant d'eau, "
                "l'aire au vent et les hauteurs des deux centres).")
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, numeric_except=("Profil", "Source"),
                              spec=SPEC_PROFILS_VENT, rejetees=rejetees)
        dicts = [d for d in dicts if str(d.get("Profil", "")).strip()]
        if not dicts:
            raise ValueError("Aucune ligne exploitable.")
        dicts.sort(key=lambda r: (str(r.get("Profil")),
                                  float(r.get("TE_milieu_m", 0.0))))
        self.profils_vent_rows = dicts
        profils = sorted({str(d["Profil"]).strip() for d in dicts})
        return {"lignes": len(dicts), "ignorees": ignored,
                "rejetees": rejetees, "profils": profils}

    def profils_vent_declares(self):
        """Les identifiants de configuration présents dans la table."""
        vus, out = set(), []
        for r in self.profils_vent_rows:
            pid = str(r.get("Profil") or "").strip()
            if pid and pid not in vus:
                vus.add(pid)
                out.append(pid)
        return out

    # ------------------------------------------------------- couples
    def import_couples(self, path):
        """Importe la table des couples (n, x_m) depuis un CSV ou un classeur."""
        headers, rows = read_table(path)
        return self.charger_couples(headers, rows)

    def charger_couples(self, headers, rows):
        mapping, ignored, missing = match_columns(headers, SPEC_COUPLES)
        if missing:
            raise ValueError(
                "Colonnes obligatoires introuvables : " + ", ".join(missing)
                + " (attendu : le numéro de couple et son abscisse).")
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, spec=SPEC_COUPLES, rejetees=rejetees)
        paires = []
        for d in dicts:
            n = d.get("n")
            if n is None or abs(n - round(n)) > 1e-6 or n < 0:
                rejetees.append(f"numéro de couple {n!r}")
                continue
            paires.append((int(round(n)), float(d["x_m"])))
        if not paires:
            raise ValueError("Aucune ligne de couple exploitable.")
        self.set_couples(paires)
        return {"couples": len(paires), "ignorees": ignored,
                "rejetees": rejetees,
                "x": (min(x for _, x in paires), max(x for _, x in paires))}

    def set_couples(self, paires):
        """Range la table des couples depuis des (n, x) : une liste indexée par
        n, `nan` là où le dossier ne donne pas de couple — c'est la forme que
        lit le calage des plans (`app_paths.table_couples`)."""
        paires = [(int(n), float(x)) for n, x in paires]
        if not paires:
            self.couples = []
            return
        nmax = max(n for n, _ in paires)
        table = [float("nan")] * (nmax + 1)
        for n, x in paires:
            table[n] = x
        self.couples = table

    def couples_paires(self):
        """[(n, x)] des seuls couples renseignés, dans l'ordre des n."""
        return [(n, x) for n, x in enumerate(self.couples) if x == x]

    def generer_couples(self, x_c0, troncons):
        """Table des couples déduite des espacements du plan de structure.

        `troncons` : [(n_debut, n_fin, espacement_m), …] — « de C.0 à C.14,
        tous les 1,00 m ». `x_c0` est l'abscisse du premier couple du premier
        tronçon (C.0 en général).

        Les tronçons doivent s'ENCHAÎNER — la fin de l'un est le début du
        suivant. Un trou laisserait des couples sans abscisse, et le calage
        d'un plan proposerait alors « sur le couple 35 » avec un X faux, ce qui
        ne se voit pas à l'écran. On refuse plutôt que de combler.
        """
        if not troncons:
            raise ValueError("Aucun tronçon : donnez au moins un intervalle "
                             "de couples et son espacement.")
        attendu = None
        for i, (n1, n2, pas) in enumerate(troncons):
            n1, n2, pas = int(n1), int(n2), float(pas)
            if n2 <= n1:
                raise ValueError(f"Tronçon {i + 1} : le couple de fin ({n2}) "
                                 f"doit être après celui de début ({n1}).")
            if pas <= 0:
                raise ValueError(f"Tronçon {i + 1} : espacement nul ou négatif.")
            if attendu is not None and n1 != attendu:
                raise ValueError(
                    f"Tronçon {i + 1} : il commence au couple {n1} alors que le "
                    f"précédent finit au couple {attendu} — les tronçons "
                    "doivent s'enchaîner sans trou ni recouvrement.")
            attendu = n2
        paires = []
        x = float(x_c0)
        n_courant = int(troncons[0][0])
        paires.append((n_courant, x))
        for n1, n2, pas in troncons:
            for n in range(int(n1) + 1, int(n2) + 1):
                x += float(pas)
                paires.append((n, x))
        self.set_couples(paires)
        return {"couples": len(paires),
                "x": (paires[0][1], paires[-1][1])}

    # ------------------------------------------------- cas de référence
    def import_cas_reference(self, path):
        """Importe les cas de chargement du dossier approuvé et leurs résultats."""
        headers, rows = read_table(path)
        return self.charger_cas_reference(headers, rows)

    def charger_cas_reference(self, headers, rows):
        mapping, ignored, missing = match_columns(headers, SPEC_CAS_REFERENCE)
        if missing:
            raise ValueError("Colonnes obligatoires introuvables : "
                             + ", ".join(missing))
        rejetees = []
        dicts = rows_to_dicts(rows, mapping, numeric_except=("Cas", "Titre"),
                              spec=SPEC_CAS_REFERENCE, rejetees=rejetees)
        cas = []
        for d in dicts:
            if not str(d.get("Cas", "")).strip():
                rejetees.append("cas sans nom")
                continue
            if d.get("VCG_solide_m") is None and d.get("VCG_corrige_m") is None:
                # sans l'un des deux VCG, le cas n'est pas calculable : le dire
                # vaut mieux que de le rejouer avec un KG nul
                rejetees.append(f"{d['Cas']} : ni VCG solide ni VCG corrigé")
                continue
            cas.append(d)
        if not cas:
            raise ValueError("Aucun cas exploitable.")
        self.cas_reference = cas
        self.rejeu_cas = None          # les résultats affichés ne valent plus
        return {"cas": len(cas), "ignorees": ignored, "rejetees": rejetees}

    @staticmethod
    def vcg_solide_du_cas(cas):
        """VCG SOLIDE d'un cas, tel que l'attend `stability.gz_curve`.

        Les dossiers impriment le plus souvent le VCG déjà corrigé de la carène
        liquide (c'est ce que fait le LOCOPIAS du bord) ; le moteur, lui,
        applique la correction lui-même à partir du FSM. Compter le FSM deux
        fois abaisserait le GM d'autant sans que rien ne le signale.
        """
        solide = cas.get("VCG_solide_m")
        if solide is not None:
            return float(solide)
        poids = float(cas.get("Poids_total_t") or 0.0)
        fsm = float(cas.get("FSM_tm") or 0.0)
        return float(cas["VCG_corrige_m"]) - (fsm / poids if poids else 0.0)

    def rejouer_cas_reference(self, dossier_navire=None):
        """Recalcule chaque cas de référence et le compare au dossier approuvé.

        Rend une liste de dicts : {cas, titre, lignes, erreur}, où `lignes` est
        la liste des (grandeur, calculé, attendu, écart, tolérance, ok) — seules
        les grandeurs effectivement fournies sont comparées.

        Le moteur ne sait lire qu'un dossier navire écrit sur le disque : sans
        `dossier_navire`, le brouillon est enregistré dans un dossier temporaire
        le temps du calcul, puis celui-ci est effacé. L'utilisateur n'a donc pas
        à enregistrer son navire avant de contrôler ses tables.
        """
        if not self.cas_reference:
            self.rejeu_cas = []
            return []
        import shutil
        import tempfile
        temporaire = None
        dossier = dossier_navire
        if not dossier:
            temporaire = tempfile.mkdtemp(prefix="carene-rejeu-")
            dossier = os.path.join(temporaire, "navire")
            self.save(dossier)
        try:
            from . import hydrostatics, stability
            from .navire import Navire
            nav = Navire.load(dossier)
            out = []
            for cas in self.cas_reference:
                code = str(cas.get("Cas", "")).strip()
                titre = str(cas.get("Titre", "") or "")
                try:
                    eq = hydrostatics.solve_equilibrium(
                        nav, float(cas["Poids_total_t"]), float(cas["LCG_m"]))
                    gz = stability.gz_curve(
                        nav, eq, self.vcg_solide_du_cas(cas),
                        float(cas.get("TCG_m") or 0.0),
                        float(cas.get("FSM_tm") or 0.0))
                except Exception as e:
                    out.append({"cas": code, "titre": titre, "lignes": [],
                                "erreur": str(e)})
                    continue
                calcule = {
                    "TE_milieu_m": eq.draft_m,
                    "Assiette_m": eq.trim_m,
                    "GM_corrige_m": gz.gm_corrige_m,
                    "GZmax_m": gz.gz_max_m,
                    "Angle_GZmax_deg": gz.angle_gz_max_deg,
                }
                lignes = []
                for cle, libelle, tol_cle in GRANDEURS_CAS:
                    attendu = cas.get(cle)
                    if attendu is None:
                        continue
                    tol = TOLERANCES_CAS[tol_cle]
                    val = float(calcule[cle])
                    ecart = val - float(attendu)
                    lignes.append((libelle, val, float(attendu), ecart, tol,
                                   abs(ecart) <= tol))
                out.append({"cas": code, "titre": titre, "lignes": lignes,
                            "erreur": ""})
            self.rejeu_cas = out
            return out
        finally:
            if temporaire:
                shutil.rmtree(temporaire, ignore_errors=True)

    @staticmethod
    def ecarts_du_rejeu(resultats):
        """(nombre de comparaisons, nombre hors tolérance) d'un rejeu."""
        total = ecarts = 0
        for r in resultats or []:
            if r.get("erreur"):
                ecarts += 1
            for _lbl, _c, _a, _e, _t, ok in r.get("lignes", []):
                total += 1
                if not ok:
                    ecarts += 1
        return total, ecarts

    # ---------------------------------------------------------- classeur
    def ecrire_classeur_type(self, chemin_xlsx):
        """Écrit le classeur Excel VIDE à remplir (une feuille par table)."""
        from .classeur_navire import ecrire_classeur
        return ecrire_classeur(chemin_xlsx, draft=None)

    def ecrire_classeur_rempli(self, chemin_xlsx):
        """Le même classeur, rempli des données de ce brouillon (c'est ainsi
        qu'est produit le classeur d'exemple)."""
        from .classeur_navire import ecrire_classeur
        return ecrire_classeur(chemin_xlsx, draft=self)

    def ecrire_modeles_csv(self, dossier):
        """Les mêmes gabarits en CSV, pour qui n'a pas Excel."""
        from .classeur_navire import ecrire_modeles_csv
        return ecrire_modeles_csv(dossier, draft=self)

    def importer_classeur(self, chemin_xlsx):
        """Lit toutes les feuilles d'un classeur rempli. Rend le compte rendu
        feuille par feuille : [(feuille, statut, message)]."""
        from .classeur_navire import importer_classeur
        return importer_classeur(self, chemin_xlsx)

    # ------------------------------------------------------------ état
    @property
    def trims_hydro(self):
        return sorted({round(r["Assiette_m"], 4) for r in self.hydro_rows})

    @property
    def trims_kn(self):
        return sorted({round(r["Assiette_m"], 4) for r in self.kn_rows})

    def _draft_range(self):
        if not self.hydro_rows:
            return None
        te = [r["TE_milieu_m"] for r in self.hydro_rows]
        return min(te), max(te)

    def ready_to_compute(self):
        """L'essentiel est-il là pour calculer un équilibre et une courbe GZ ?

        Une seule assiette hydrostatique suffit au moteur (l'assiette
        d'équilibre n'est alors pas interpolée — l'étape le signale) : exiger
        deux assiettes ici interdirait un navire que le moteur sait calculer."""
        return (len(self.trims_hydro) >= 1 and len(self.hydro_rows) >= 4
                and len(self.trims_kn) >= 1 and len(self.kn_rows) >= 2
                and bool(self.kn_angles)
                and all(k in self.lege for k in ("masse_t", "lcg_m", "vcg_m"))
                and bool(self.identification.get("nom")))

    def status(self, project=None):
        """{clé d'étape: (statut, [(texte, ton), ...])} pour toute la fenêtre."""
        st = {}

        # --- le classeur à remplir : point d'entrée, jamais un passage obligé
        if self.classeur_importe:
            st["classeur"] = ("ok", [
                (f"Classeur importé : {self.classeur_importe}.", "ok"),
                ("Les étapes suivantes montrent ce qui en a été lu.", "ok")])
        else:
            st["classeur"] = ("info", [
                ("Aucun classeur importé — vous pouvez aussi remplir les "
                 "étapes une par une.", "info")])

        # --- identification
        manque = [lbl for k, lbl, _, req in CHAMPS_IDENTIFICATION + CHAMPS_DIMENSIONS
                  if req and not self.identification.get(k)
                  and not self.dimensions.get(k)]
        conv = str(self.identification.get("convention_fsm") or "").strip().lower()
        conv_mauvaise = bool(conv) and conv not in CONVENTIONS_FSM
        if not self.identification.get("nom"):
            st["identification"] = ("todo", [("Navire non nommé.", "todo")])
        elif manque or conv_mauvaise:
            lignes = [(f"Nom : {self.identification['nom']}.", "ok")]
            if manque:
                lignes.append(("Manque : " + ", ".join(manque) + ".", "warn"))
            if conv_mauvaise:
                lignes.append((f"Carène liquide « {conv} » inconnue : écrivez "
                               "« reel » ou « max ».", "warn"))
            st["identification"] = ("warn", lignes)
        else:
            lignes = [(f"« {self.identification['nom']} » — dimensions renseignées.", "ok")]
            if not conv:
                lignes.append(("Carène liquide non précisée : le moteur prendra "
                               "le FSM maximal (« max »), le plus conservatif.",
                               "info"))
            st["identification"] = ("ok", lignes)

        # --- hydrostatiques
        trims = self.trims_hydro
        if not trims:
            st["hydro"] = ("todo", [("Aucune table importée.", "todo")])
        else:
            lines = [(f"{len(trims)} assiette(s) : "
                      + ", ".join(f"{t:+.3f} m" for t in trims), "ok"),
                     (f"{len(self.hydro_rows)} lignes au total.", "ok")]
            rng = self._draft_range()
            if rng:
                lines.append((f"Tirants d'eau {rng[0]:.2f} m → {rng[1]:.2f} m.", "ok"))
            if len(trims) < 2:
                lines.append(("Une seule assiette : l'assiette d'équilibre ne "
                              "pourra pas être interpolée.", "warn"))
                st["hydro"] = ("warn", lines)
            else:
                st["hydro"] = ("ok", lines)

        # --- pantocarènes
        if not self.kn_rows:
            st["kn"] = ("todo", [("Aucune pantocarène importée.", "todo")])
        else:
            lines = [(f"{len(self.trims_kn)} assiette(s), "
                      f"{len(self.kn_rows)} lignes.", "ok"),
                     ("Angles : " + ", ".join(f"{a}°" for a in self.kn_angles), "ok")]
            statut = "ok"
            if not self.kn_angles or max(self.kn_angles) < 40:
                lines.append(("Les angles ne vont pas jusqu'à 40° : les critères "
                              "d'aire IS2008 ne pourront pas être évalués.", "warn"))
                statut = "warn"
            st["kn"] = (statut, lines)

        # --- lège
        manque = [lbl for k, lbl, _, req in CHAMPS_LEGE
                  if req and self.lege.get(k) is None]
        if manque:
            # « à faire » tant que RIEN d'obligatoire n'est saisi ; « à
            # vérifier » dès qu'une partie l'est. Comparer à toutes les cases,
            # TCG facultatif compris, marquait un navire vierge « à vérifier »
            requis = sum(1 for _k, _l, _u, req in CHAMPS_LEGE if req)
            st["lege"] = ("todo" if len(manque) == requis else "warn",
                          [("Manque : " + ", ".join(manque) + ".", "todo")])
        else:
            st["lege"] = ("ok", [
                (f"{self.lege['masse_t']:.2f} t · LCG {self.lege['lcg_m']:.3f} m · "
                 f"VCG {self.lege['vcg_m']:.3f} m", "ok")])

        # --- critères : les réglementations retenues
        ids = self.reglements_retenus_ids()
        n = len(ids)
        lignes = [(f"{n} réglementation(s) retenue(s).", "ok" if n else "warn")]
        if not self.profil.get("type"):
            lignes.append(("Profil du navire non renseigné : les propositions de la "
                           "bibliothèque ne le connaissent pas.", "warn"))
        st["criteres"] = ("ok" if n else "warn", lignes)

        # --- capacités
        if not self.capacites:
            st["capacites"] = ("todo", [("Aucune capacité déclarée.", "todo")])
        else:
            sans = [c["Nom"] for c in self.capacites if c["Nom"] not in self.jauges]
            lines = [(f"{len(self.capacites)} capacité(s) déclarée(s), "
                      f"{len(self.jauges)} table(s) de jaugeage.", "ok")]
            if sans:
                lines.append((f"Sans jaugeage : {', '.join(sans[:6])}"
                              + (" …" if len(sans) > 6 else ""), "warn"))
            st["capacites"] = ("warn" if sans else "ok", lines)

        # --- envahissement
        if not self.envahissement:
            st["envahissement"] = ("todo", [
                ("Aucun point — l'angle d'envahissement ne sera pas calculé.",
                 "todo")])
        else:
            sans_desc = sum(1 for p in self.envahissement if not p.get("Fonction"))
            lines = [(f"{len(self.envahissement)} point(s) saisi(s).", "ok")]
            if sans_desc:
                lines.append((f"{sans_desc} point(s) sans description.", "warn"))
            st["envahissement"] = ("warn" if sans_desc else "ok", lines)

        # --- vent (étape entièrement facultative : ce n'est pas une donnée du
        # navire mais d'une configuration de référence)
        renseignes = [k for k, _l, _u, _r in CHAMPS_VENT
                      if self.vent.get(k) is not None]
        if len(renseignes) == len(CHAMPS_VENT):
            st["vent"] = ("ok", [
                (f"Surface au vent {self.vent['aire_windage_m2']:.1f} m² "
                 f"à {self.vent['vcw_m']:.2f} m (configuration de "
                 "référence).", "ok")])
        elif renseignes:
            st["vent"] = ("warn", [
                ("Renseignement partiel : les quatre valeurs sont nécessaires "
                 "au critère météo.", "warn")])
        else:
            st["vent"] = ("info", [
                ("Facultatif — sans ces valeurs, le critère météo n'est "
                 "simplement pas évalué.", "info")])
        # la table par voilure, quand elle est là, dit ce que les quatre
        # champs ne peuvent pas dire : de quelle toile on parle
        if self.profils_vent_rows:
            profils = self.profils_vent_declares()
            st["vent"] = ("ok", list(st["vent"][1]) + [
                (f"Surfaces au vent par voilure : {len(profils)} "
                 f"configuration(s) ({', '.join(profils)}), "
                 f"{len(self.profils_vent_rows)} mesures.", "ok")])

        # --- couples (calage des plans)
        paires = self.couples_paires()
        if not paires:
            st["couples"] = ("todo", [
                ("Aucun couple — le calage des plans se fera en mesurant les "
                 "abscisses à la main.", "todo")])
        else:
            xs = [x for _n, x in paires]
            lines = [(f"{len(paires)} couples, de {xs[0]:.3f} à {xs[-1]:.3f} m "
                      f"(C.{paires[0][0]} à C.{paires[-1][0]}).", "ok")]
            croissants = all(b > a for a, b in zip(xs, xs[1:]))
            if not croissants:
                lines.append(("Les abscisses ne sont pas croissantes : deux "
                              "couples se croisent, ou un signe est inversé.",
                              "warn"))
            st["couples"] = ("ok" if croissants else "warn", lines)

        st.update(self._status_plans(project))

        # --- cas de référence (contrôle)
        if not self.cas_reference:
            st["validation"] = ("todo", [
                ("Aucun cas de référence — contrôle non effectué.", "todo")])
        elif self.rejeu_cas is None:
            st["validation"] = ("warn", [
                (f"{len(self.cas_reference)} cas saisis, pas encore rejoués.",
                 "warn")])
        else:
            total, ecarts = self.ecarts_du_rejeu(self.rejeu_cas)
            if ecarts:
                st["validation"] = ("warn", [
                    (f"{ecarts} écart(s) hors tolérance sur {total} "
                     f"comparaison(s), {len(self.rejeu_cas)} cas rejoués.",
                     "warn"),
                    ("Un écart signale une table mal reprise : reprenez la "
                     "colonne en cause avant de mettre le navire en service.",
                     "warn")])
            else:
                st["validation"] = ("ok", [
                    (f"{total} comparaison(s) dans la tolérance sur "
                     f"{len(self.rejeu_cas)} cas : les tables sont bonnes.",
                     "ok")])
        pret = self.ready_to_compute()
        lines = [("Prêt à calculer : oui — l'essentiel est complet." if pret
                  else "Prêt à calculer : non — il manque des données "
                       "essentielles.", "ok" if pret else "todo")]
        for key in ("capacites", "envahissement", "vent"):
            if st.get(key, ("ok", []))[0] != "ok":
                lines.append((f"{key.capitalize()} : à compléter.", "warn"))
        st["recap"] = ("ok" if pret else "todo", lines)
        # --- LA COHÉRENCE DES TABLES (D-82) : ce que la physique impose aux
        # tables, relu à chaque état de la fenêtre. Une erreur fait passer
        # l'étape en « à vérifier », quel que soit le reste ; les constats
        # s'ajoutent aux lignes de l'étape qu'ils concernent.
        try:
            from . import controles_tables as _C
            constats = _C.controler_brouillon(self)
        except Exception:                       # noqa: BLE001 — jamais bloquant
            constats = []
        self.constats_tables = constats
        for table, cle in (("Hydrostatiques", "hydro"), ("Pantocarènes", "kn"),
                           ("Jaugeages", "capacites")):
            propres = [c for c in constats if c.table == table and c.niveau != _C.INFO]
            if not propres or cle not in st:
                continue
            statut, lignes = st[cle]
            lignes = list(lignes) + [(c.message, "warn") for c in propres[:4]]
            if len(propres) > 4:
                lignes.append((f"… et {len(propres) - 4} autre(s) : voir « Contrôle ».",
                               "warn"))
            if any(c.niveau == _C.ERREUR for c in propres) or statut == "ok":
                statut = "warn"
            st[cle] = (statut, lignes)
        return st

    @staticmethod
    def _status_plans(project):
        """État des étapes « plans », déduit du fichier .carene.json lié."""
        if project is None:
            msg = [("Aucun fichier de plans lié (facultatif).", "todo")]
            return {k: ("todo", msg) for k in
                    ("profil", "ponts", "plans_ponts", "polygones")}
        out = {}
        out["profil"] = (("ok", [("Profil importé et calé.", "ok")])
                         if project.profile.ready else
                         ("todo", [("Profil non calé.", "todo")]))
        n_ponts = len(project.decks)
        out["ponts"] = (("ok", [(f"{n_ponts} pont(s) défini(s).", "ok")])
                        if n_ponts else ("todo", [("Aucun pont.", "todo")]))
        cales = [d for d in project.decks if not d.plan.ready]
        if not n_ponts:
            out["plans_ponts"] = ("todo", [("Aucun pont.", "todo")])
        elif cales:
            out["plans_ponts"] = ("warn", [
                (f"{n_ponts - len(cales)}/{n_ponts} plan(s) calé(s).", "warn")])
        else:
            out["plans_ponts"] = ("ok", [(f"{n_ponts} plan(s) calé(s).", "ok")])
        from ..project import KIND_CONTOUR
        caps = [c for _, c in project.all_capacities() if c.kind != KIND_CONTOUR]
        out["polygones"] = (("ok", [(f"{len(caps)} capacité(s) tracée(s).", "ok")])
                            if caps else
                            ("todo", [("Aucune capacité tracée.", "todo")]))
        return out

    # ------------------------------------------------- les réglementations
    def reglements_retenus_ids(self):
        """Les identifiants des réglementations retenues : le choix fait dans
        la bibliothèque, ou, faute de choix, la traduction des anciens jeux de
        critères (et la ligne de charge)."""
        from . import reglements as R
        if self.reglements_actifs is not None:
            return [rid for rid, on in self.reglements_actifs.items() if on]
        ids = [R.CORRESPONDANCE_ANCIENNE[k][0] for k, on in self.criteres_actifs.items()
               if on and k in R.CORRESPONDANCE_ANCIENNE]
        return ids + ["LL1966_LIGNE_DE_CHARGE"]

    def proposer_reglements(self):
        """Coche ce que le profil appelle, et rien d'autre. Rend les
        identifiants proposés."""
        from . import reglements as R
        man = {"profil": self.profil, "dimensions": self.dimensions}
        ids = [r.id for r, _raison in R.proposer(R.profil_du_navire(man))]
        regs, _m = R.bibliotheque()
        self.reglements_actifs = {rid: (rid in ids) for rid in regs}
        self._synchroniser_criteres()
        return ids

    def choisir_reglement(self, rid, actif):
        from . import reglements as R
        if self.reglements_actifs is None:
            self.reglements_actifs = {x: True for x in self.reglements_retenus_ids()}
            for x in R.bibliotheque()[0]:
                self.reglements_actifs.setdefault(x, False)
        self.reglements_actifs[rid] = bool(actif)
        self._synchroniser_criteres()

    def _synchroniser_criteres(self):
        """Les anciens jeux de critères (`criteres`, lus par les Carène 2.x)
        suivent le choix de la bibliothèque : un navire 3.0 s'ouvre encore
        dans une 2.x avec les mêmes critères."""
        from . import reglements as R
        if self.reglements_actifs is None:
            return
        for cle, (rid, _cles) in R.CORRESPONDANCE_ANCIENNE.items():
            self.criteres_actifs[cle] = bool(self.reglements_actifs.get(rid))

    def _reglements_pour_manifeste(self):
        """La liste `navire.json › reglements`, surcharges du dossier
        comprises (seuils et conventions lus dans les anciens `criteres`)."""
        from . import reglements as R

        class _N:
            pass
        faux = _N()
        faux.manifest = {"criteres": {k: v for k, v in self.criteres.items()}}
        faux.criteres = faux.manifest["criteres"]
        anciennes, _m = R.reglements_du_navire(faux)
        par_id = {x.reglement.id: x for x in anciennes}
        regs, _m = R.bibliotheque()
        retenues = []
        for rid in self.reglements_retenus_ids():
            if rid in par_id:
                retenues.append(par_id[rid])
            elif rid in regs:
                retenues.append(R.Retenue(regs[rid]))
        return R.pour_navire_json(retenues)

    # ------------------------------------------------------------ écriture
    def _manifest(self):
        man = dict(self._extra_manifest)
        man.update({
            "format_version": 1,
            "identification": self.identification,
            "dimensions": self.dimensions,
            "reperes": self.reperes,
            "lege": self.lege,
            "criteres": {k: v for k, v in self.criteres.items()
                         if self.criteres_actifs.get(k, True)},
            # 3.0 (D-79) : les réglementations retenues, avec leur version
            "reglements": self._reglements_pour_manifeste(),
        })
        if self.profil:
            man["profil"] = dict(self.profil)
        # perpendiculaires : écrites dans la structure que lit le moteur
        # (stability.abscisses_tirants_eau) dès que les deux X sont saisis
        x_ar = self.dimensions.get("x_perpendiculaire_ar_m")
        x_av = self.dimensions.get("x_perpendiculaire_av_m")
        if x_ar is not None and x_av is not None:
            man["dimensions"] = dict(man["dimensions"])
            man["dimensions"]["perpendiculaires"] = {
                "arriere_m": float(x_ar), "avant_m": float(x_av)}
        if self.vent:
            man["profil_vent_reference"] = self.vent
        # La LISTE des configurations de voilure accompagne la table : c'est
        # elle qui donne les libellés du sélecteur à l'écran. Celles que le
        # dossier décrit déjà (nom, description) sont gardées telles quelles ;
        # une configuration mesurée mais non décrite est déclarée sous son
        # identifiant, faute de mieux — mieux vaut un libellé pauvre qu'une
        # table inutilisable.
        declares = list(man.get("profils_vent") or [])
        connus = {str(p.get("id")) for p in declares if p.get("id")}
        for pid in self.profils_vent_declares():
            if pid not in connus:
                declares.append({"id": pid, "nom": pid})
        if declares:
            man["profils_vent"] = declares
        # la convention de carène liquide est saisie avec l'identification
        # mais vit en tête du manifeste, où `Navire.load` la lit
        ident = dict(self.identification)
        conv = str(ident.pop("convention_fsm", "") or "").strip().lower()
        man["identification"] = ident
        if conv in CONVENTIONS_FSM:
            man["convention_fsm"] = conv
        man["genere_par"] = "Carène — fenêtre « Création du navire »"
        return man

    @staticmethod
    def _colonnes(rows, premieres):
        """Union ordonnée des colonnes de toutes les lignes, `premieres` en
        tête. Prendre les clés de la première ligne seule perdait les colonnes
        facultatives absentes de cette ligne mais présentes ailleurs (import
        d'un fichier à cellules vides, ou de deux fichiers à colonnes
        différentes). Les cellules manquantes sont écrites "" ; le chargeur
        (`navire._grid2d_from_rows`, `Capacity`) sait les ignorer."""
        cols = list(premieres)
        vus = set(cols)
        for r in rows:
            for c in r:
                if c not in vus:
                    vus.add(c)
                    cols.append(c)
        return cols

    @staticmethod
    def _write_csv(path, rows, columns):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with ecriture_atomique(path, newline="") as f:
            w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            w.writeheader()
            for r in rows:
                w.writerow(r)

    def save(self, path):
        """Écrit le dossier navire virtuel. Retourne la liste des fichiers."""
        os.makedirs(path, exist_ok=True)
        written = []

        with ecriture_atomique(os.path.join(path, "navire.json")) as f:
            json.dump(self._manifest(), f, ensure_ascii=False, indent=1)
        written.append("navire.json")

        if self.hydro_rows:
            cols = self._colonnes(self.hydro_rows, ["Assiette_m", "TE_milieu_m"])
            self._write_csv(os.path.join(path, "hydrostatiques.csv"),
                            self.hydro_rows, cols)
            written.append("hydrostatiques.csv")

        if self.kn_rows:
            cols = ["Assiette_m", "Deplacement_t"] + [
                f"KN_{float(a):g}" for a in self.kn_angles]
            self._write_csv(os.path.join(path, "pantocarenes_kn.csv"),
                            self.kn_rows, cols)
            written.append("pantocarenes_kn.csv")

        if self.capacites:
            cols = ["Nom", "Type", "Groupe", "Densite", "Volume_net_m3",
                    "Baffles", "Perm_pc", "Poids_t", "LCG_m", "TCG_m", "VCG_m",
                    "FSM_max_tm", "Ballast"]
            rows = [{c: r.get(c, "") for c in cols} for r in self.capacites]
            self._write_csv(os.path.join(path, "capacites.csv"), rows, cols)
            written.append("capacites.csv")

        # deux noms qui donnent le même fichier (« WB 2C » / « WB/2C ») se
        # liraient l'un l'autre en silence : on refuse d'écrire un dossier
        # ambigu plutôt que de compter deux fois le même liquide
        fichiers = {}
        for nom in self.jauges:
            fn = nom.replace(" ", "_").replace("/", "_") + ".csv"
            if fn in fichiers:
                raise ValueError(
                    f"Les capacités « {fichiers[fn]} » et « {nom} » donnent le "
                    f"même fichier de jaugeage ({fn}) : renommez l'une des "
                    "deux avant d'enregistrer.")
            fichiers[fn] = nom
        for nom, rows in self.jauges.items():
            cols = self._colonnes(rows, ["Remplissage_pc"])
            fn = nom.replace(" ", "_").replace("/", "_") + ".csv"
            self._write_csv(os.path.join(path, "jauges", fn), rows, cols)
            written.append(f"jauges/{fn}")

        if self.envahissement:
            cols = ["Repere", "Fonction", "X_m", "Y_m", "Z_m"]
            self._write_csv(os.path.join(path, "points_envahissement.csv"),
                            self.envahissement, cols)
            written.append("points_envahissement.csv")

        # les surfaces au vent par voilure : un fichier du dossier comme les
        # autres, et la liste des configurations est écrite avec, dans le
        # manifeste (voir `_manifest`) — la table sans la liste donnerait un
        # sélecteur sans libellés
        if self.profils_vent_rows:
            cols = [c.name for c in SPEC_PROFILS_VENT]
            rows = [{c: r.get(c, "") for c in cols}
                    for r in self.profils_vent_rows]
            self._write_csv(os.path.join(path, "profils_vent.csv"), rows, cols)
            written.append("profils_vent.csv")

        # la table des couples : mêmes colonnes que le fichier déjà lu par
        # `app_paths.table_couples`, sans quoi le calage des plans ne la
        # reconnaîtrait pas
        paires = self.couples_paires()
        if paires:
            self._write_csv(os.path.join(path, "couples.csv"),
                            [{"n": n, "x_m": f"{x:.3f}"} for n, x in paires],
                            ["n", "x_m"])
            written.append("couples.csv")

        if self.cas_reference:
            cols = [c.name for c in SPEC_CAS_REFERENCE]
            rows = [{c: r.get(c, "") for c in cols} for r in self.cas_reference]
            self._write_csv(os.path.join(path, "validation", "cas_reference.csv"),
                            rows, cols)
            written.append("validation/cas_reference.csv")

        return written

    # ------------------------------------------------------------ lecture
    @classmethod
    def load(cls, path):
        """Recharge un dossier navire virtuel existant dans un brouillon."""
        d = cls()
        d.source_path = path
        with open(os.path.join(path, "navire.json"), encoding="utf-8") as f:
            man = json.load(f)
        d.identification = dict(man.get("identification", {}))
        if man.get("convention_fsm") and "convention_fsm" not in d.identification:
            d.identification["convention_fsm"] = man["convention_fsm"]
        d.dimensions = man.get("dimensions", {})
        # les perpendiculaires sont écrites en sous-dictionnaire (c'est la
        # forme que lit le moteur) ; le formulaire, lui, a deux champs plats.
        # Sans ce dépliage, rouvrir un dossier qui les porte montrait deux
        # cases vides — et les réenregistrer les effaçait.
        pp = d.dimensions.get("perpendiculaires") or {}
        for cle_plate, cle_nichee in (("x_perpendiculaire_ar_m", "arriere_m"),
                                      ("x_perpendiculaire_av_m", "avant_m")):
            if pp.get(cle_nichee) is not None and cle_plate not in d.dimensions:
                d.dimensions[cle_plate] = pp[cle_nichee]
        d.reperes = man.get("reperes", d.reperes)
        d.lege = man.get("lege", {})
        if man.get("criteres"):
            d.criteres = {**CRITERES_DEFAUT, **man["criteres"]}
            d.criteres_actifs = {k: (k in man["criteres"]) for k in d.criteres}
        d.vent = man.get("profil_vent_reference", {})
        d.profil = dict(man.get("profil") or {})
        if isinstance(man.get("reglements"), list):
            from . import reglements as R
            ids = {str(x.get("id")) for x in man["reglements"]}
            d.reglements_actifs = {rid: (rid in ids) for rid in R.bibliotheque()[0]}
            for rid in ids:
                d.reglements_actifs.setdefault(rid, True)
        d._extra_manifest = {k: v for k, v in man.items() if k not in (
            "format_version", "identification", "dimensions", "reperes",
            "lege", "criteres", "profil_vent_reference", "genere_par",
            "convention_fsm", "profil", "reglements")}

        def rows_of(fn):
            p = os.path.join(path, fn)
            if not os.path.exists(p):
                return []
            with open(p, newline="", encoding="utf-8") as f:
                return [{k: v for k, v in r.items()} for r in csv.DictReader(f)]

        def numify(rows, text_cols=()):
            out = []
            for r in rows:
                dd = {}
                for k, v in r.items():
                    if k in text_cols:
                        dd[k] = v
                    else:
                        fv = to_float(v)
                        dd[k] = fv if fv is not None else v
                out.append(dd)
            return out

        d.hydro_rows = numify(rows_of("hydrostatiques.csv"))
        d.kn_rows = numify(rows_of("pantocarenes_kn.csv"))
        if d.kn_rows:
            d.kn_angles = sorted(float(k.split("_")[1]) for k in d.kn_rows[0]
                                 if k.startswith("KN_"))
        d.capacites = numify(rows_of("capacites.csv"),
                             text_cols=("Nom", "Type", "Groupe"))
        d.envahissement = numify(rows_of("points_envahissement.csv"),
                                 text_cols=("Repere", "Fonction"))
        d.profils_vent_rows = [
            {k: v for k, v in r.items() if v != ""}
            for r in numify(rows_of("profils_vent.csv"),
                            text_cols=("Profil", "Source"))]
        jdir = os.path.join(path, "jauges")
        if os.path.isdir(jdir):
            for fn in sorted(os.listdir(jdir)):
                if not fn.endswith(".csv"):
                    continue
                nom = fn[:-4].replace("_", " ")
                d.jauges[nom] = numify(rows_of(os.path.join("jauges", fn)))
        # recale les noms de jauges sur ceux déclarés dans capacites.csv
        declared = {c["Nom"] for c in d.capacites}
        for nom in list(d.jauges):
            if nom not in declared:
                for cand in declared:
                    if cand.replace(" ", "_") == nom.replace(" ", "_"):
                        d.jauges[cand] = d.jauges.pop(nom)
                        break

        paires = []
        for r in rows_of("couples.csv"):
            n, x = to_float(r.get("n")), to_float(r.get("x_m"))
            if n is not None and x is not None:
                paires.append((int(round(n)), x))
        d.set_couples(paires)

        d.cas_reference = numify(rows_of(os.path.join("validation",
                                                      "cas_reference.csv")),
                                 text_cols=("Cas", "Titre"))
        # une cellule vide relue vaut "" : la clé doit DISPARAÎTRE, sinon le
        # rejeu comparerait une grandeur que le dossier n'a pas imprimée
        d.cas_reference = [{k: v for k, v in c.items() if v != ""}
                           for c in d.cas_reference]
        return d
